"""
PMT Alarm – Closed Alarm downloader (Playwright edition) – CALENDAR + PER-PAGE CYCLE EDITION

The site resets "Items per page" when the table moves to another page, so the page size can NOT be set once.
FLOW, repeated for EVERY page N:
  1. ROW SIZE   : set Items per page to 500, wait for the reload, verify the table really shows 500 rows
                  (or all remaining rows when fewer exist). Never continue with 10 rows.
  2. START DATE : pick the Start Date through the REAL calendar widget (open it, navigate month/year,
                  click the day). No typing, no fill, no keyboard entry, no JS value injection.
  3. END DATE   : NEVER touched. The site keeps its own default (today).
  4. APPLY      : click Apply, wait for loading indicators, table redraw and a stable table.
  5. VERIFY SIZE: the site may reset the size after Apply -> check again, restore 500 if needed.
  6. NAVIGATE   : for N > 1 go to page N (page-number button, or Next); verify the footer range
                  ("501-1000 of ...") and verify again that the size is still 500 (restore if it reset).
  7. DOWNLOAD   : Download Alarm > By Page (one file per page); the file must exist and be non-empty.
  8. NEXT PAGE  : repeat the whole cycle for page N+1 until the last page.
All page files are compiled into ONE Excel file.
STEP 2  Download -> hand the compiled XLSX to the browser (no processing)

MEMORY MANAGEMENT
  * Guard measures the container WORKING SET (memory.current - inactive_file), 93 %, 3 samples in a row,
    kills only Chromium.
  * downloads are deleted from Playwright's store, JS handles are disposed, light-weight page-change signature,
    fonts/images/media blocked, small JPEG screenshots, capped requestfailed logging.
  * Browser RECYCLING: on memory pressure Chromium is closed and relaunched. Because every page starts with its
    own complete cycle, resuming at page N is cheap.
  * Per-page diagnostics (memory before/after, rows inside the downloaded file vs rows on screen).

All state lives on disk in a per-session job folder (not only in st.session_state),
so it survives websocket reconnects and page refreshes (the job id is kept in the URL).
"""
import csv
import faulthandler
import hashlib
import io
import json
import logging
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import traceback
import uuid
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import streamlit as st
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter
from playwright.sync_api import Error as PWError
from playwright.sync_api import TimeoutError as PWTimeout
from playwright.sync_api import sync_playwright

try:
    import psutil
except ImportError:  # diagnostics degrade gracefully
    psutil = None

# ===================== CONFIGURATION =====================
LOGIN_URL = "https://pmt-alarm.komdigi.go.id/auth/login"
ALARM_URL = "https://pmt-alarm.komdigi.go.id/dashboard/alarm"
DEFAULT_TIMEOUT = 40          # seconds for normal waits
JOBS_ROOT = Path(tempfile.gettempdir()) / "pmt_alarm_jobs"
STALE_AFTER = 600             # a "running" step with no heartbeat for 10 min = interrupted
JOB_MAX_AGE_HOURS = 12
POLL_SECONDS = 3
MONITOR_INTERVAL = 2          # seconds between diagnostic samples
BLOCK_IMAGES = True           # also blocks media + fonts
MEM_LIMIT_MB_OVERRIDE = float(os.environ.get("MEM_LIMIT_MB", 0)) or None

# ---- per-page cycle -----------------------------------------------------
TARGET_PAGE_SIZE = 500        # the ONLY page size used (never derived from the rows visible on screen)
SIZE_VERIFY_WAIT = 10         # seconds to wait for the table to reach the wanted size before re-selecting it
SIZE_RESTORE_ATTEMPTS = 4     # how many times a reset page size is restored before the cycle fails
STABLE_SECONDS = 2.0          # the table signature must stay unchanged this long after Apply
CAL_MAX_MONTH_CLICKS = 60     # safety cap when navigating months in the calendar
APPLY_QUICK_TIMEOUT = 10      # seconds to wait for a visible change after Apply (cycles after the first one)
MAX_NAV_STEPS = 120           # max clicks to reach page N
CYCLE_RETRIES = 2             # a failed page cycle is retried this many times inside the same browser session

# ---- memory management -------------------------------------------------
MEM_GUARD_FRACTION = 0.93     # HARD guard: working set above 93 % of the limit ...
MEM_GUARD_SAMPLES = 3         # ... for this many consecutive samples (x MONITOR_INTERVAL s) kills Chromium
SOFT_RECYCLE_FRACTION = 0.80  # SOFT limit: after a page, recycle the browser if working set is above this
CHROMIUM_GROWTH_MB = 450      # ... or if Chromium RSS grew this much since the first page of the session
RECYCLE_MIN_PAGES = 3         # never recycle before this many pages were downloaded in a session
RECYCLE_EVERY_PAGES = 0       # >0 = also recycle the browser every N pages (0 = only on memory pressure)
MAX_RECYCLES = 80             # planned recycles (memory pressure) allowed per job
MAX_CRASH_RESTARTS = 6        # unplanned restarts (crash / guard trip / error) that made progress
PROBE_FIRST_PAGES = 3         # verify the real row count inside the first N downloaded files ...
PROBE_EVERY = 25              # ... and then every N pages

SITE_UTC_OFFSET_HOURS = 7     # the site works in WIB (UTC+7)
DATE_INPUT_FORMAT = "%d-%m-%Y"        # how the Start Date box displays dates (02-10-2026)
PAGE_SIZE_CHOICES = ("500", "250", "100", "50")  # (kept for reference; the flow uses TARGET_PAGE_SIZE only)
DEFAULT_PAGE_ROWS = 10        # rows the site shows before a bigger page size is chosen
MAX_PAGES = 3000              # safety cap for the pagination loop
DOWNLOAD_TIMEOUT = 240        # seconds to wait for ONE "By Page" file
PAGE_SIZE_LOAD_TIMEOUT = 240  # seconds to wait for the table to grow after choosing a bigger page size
PAGE_SIZE_STABLE_SECONDS = 6  # a partial table (fewer rows than chosen size) must stay unchanged this long
APPLY_LOAD_TIMEOUT = 90       # seconds to wait for the table to reload after the FIRST Apply of a session

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
# =========================================================

WIB = timezone(timedelta(hours=SITE_UTC_OFFSET_HOURS))
QUIET_PHASES = ("finished", "closing", "recycling", "compile", "idle")

logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
log = logging.getLogger("pmt_alarm")
faulthandler.enable()

_INSTALL_LOCK = threading.Lock()


# =====================================================================
#  Job folder, status files and log (shared by the UI and the workers)
# =====================================================================
def write_json_atomic(path: Path, data: dict):
    tmp = path.with_name(path.name + ".tmp")
    tmp.write_text(json.dumps(data), encoding="utf-8")
    os.replace(tmp, path)


def job_log(job_dir: Path, msg: str, level: str = "INFO"):
    line = f"{datetime.now():%H:%M:%S} [{level}] {msg}"
    try:
        with open(job_dir / "job.log", "a", encoding="utf-8") as fh:
            fh.write(line + "\n")
    except OSError:
        pass
    log.log(getattr(logging, level, logging.INFO), "[%s] %s", job_dir.name, msg)


class StepReporter:
    """Writes step1.json so the UI can poll progress."""

    def __init__(self, job_dir: Path, step: int):
        self.job_dir, self.step, self.t0 = job_dir, step, time.time()

    def _write(self, state, message, progress=None, **extra):
        write_json_atomic(
            self.job_dir / f"step{self.step}.json",
            {"state": state, "message": message, "progress": progress,
             "started": self.t0, "updated": time.time(), **extra},
        )

    def running(self, message, progress=None, quiet=False):
        if not quiet:
            job_log(self.job_dir, f"[step {self.step}] {message}")
        self._write("running", message, progress)

    def done(self, message, **extra):
        job_log(self.job_dir, f"[step {self.step}] DONE: {message}")
        self._write("done", message, 1.0, **extra)

    def fail(self, message):
        job_log(self.job_dir, f"[step {self.step}] FAILED: {message}", "ERROR")
        self._write("error", message)


def get_status(job_dir: Path, step: int) -> dict:
    path = job_dir / f"step{step}.json"
    if not path.exists():
        return {"state": "idle", "message": ""}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {"state": "idle", "message": ""}
    if data.get("state") == "running" and time.time() - data.get("updated", 0) > STALE_AFTER:
        data["state"] = "error"
        data["message"] = "This step was interrupted (the app may have restarted). Please run it again."
    return data


def cleanup_old_jobs():
    if not JOBS_ROOT.exists():
        return
    cutoff = time.time() - JOB_MAX_AGE_HOURS * 3600
    for d in JOBS_ROOT.iterdir():
        try:
            if d.is_dir() and d.stat().st_mtime < cutoff:
                shutil.rmtree(d, ignore_errors=True)
        except OSError:
            pass


# =====================================================================
#  Diagnostics: container memory, process tree, disk, browser events
# =====================================================================
def _read_int(path):
    try:
        raw = Path(path).read_text().strip()
        return None if raw == "max" else int(raw)
    except Exception:
        return None


def _mb(v):
    return None if v is None else v / 1024 / 1024


def _cg_stat() -> dict:
    """memory.stat of the container cgroup (v2 or v1)."""
    for p in ("/sys/fs/cgroup/memory.stat", "/sys/fs/cgroup/memory/memory.stat"):
        try:
            out = {}
            for line in Path(p).read_text().splitlines():
                parts = line.split()
                if len(parts) == 2:
                    try:
                        out[parts[0]] = int(parts[1])
                    except ValueError:
                        pass
            if out:
                return out
        except Exception:
            continue
    return {}


def cgroup_memory() -> dict:
    used = _read_int("/sys/fs/cgroup/memory.current")
    limit = _read_int("/sys/fs/cgroup/memory.max")
    peak = _read_int("/sys/fs/cgroup/memory.peak")
    if used is None:  # cgroup v1
        used = _read_int("/sys/fs/cgroup/memory/memory.usage_in_bytes")
        limit = _read_int("/sys/fs/cgroup/memory/memory.limit_in_bytes")
        peak = _read_int("/sys/fs/cgroup/memory/memory.max_usage_in_bytes")
    if limit is not None and limit > (1 << 40):
        limit = None
    oom = None
    try:
        for line in Path("/sys/fs/cgroup/memory.events").read_text().splitlines():
            if line.startswith("oom_kill "):
                oom = int(line.split()[1])
    except Exception:
        pass
    st_ = _cg_stat()
    inactive = st_.get("inactive_file", st_.get("total_inactive_file"))
    # working set = what the kernel can NOT simply drop (this is what really leads to an OOM kill)
    ws = used - inactive if (used is not None and inactive is not None) else used
    return {
        "used_mb": _mb(used), "ws_mb": _mb(ws),
        "limit_mb": MEM_LIMIT_MB_OVERRIDE or _mb(limit),
        "peak_mb": _mb(peak), "oom_kills": oom,
        "anon_mb": _mb(st_.get("anon", st_.get("total_rss"))),
        "file_mb": _mb(st_.get("file", st_.get("total_cache"))),
        "shmem_mb": _mb(st_.get("shmem", st_.get("total_shmem"))),
    }


def process_tree():
    out = []
    if psutil is None:
        return out
    try:
        me = psutil.Process(os.getpid())
        for p in [me] + me.children(recursive=True):
            try:
                cmd = " ".join(p.cmdline())
                name = p.name()
                m = re.search(r"--type=([\w-]+)", cmd)
                if p.pid == me.pid:
                    role = "streamlit"
                elif name.startswith("node"):
                    role = "playwright-node"
                elif "chrom" in name.lower() or "chrom" in cmd.lower():
                    role = "chromium-" + (m.group(1) if m else "browser")
                else:
                    role = name
                out.append((p.pid, role, p.status(), p.memory_info().rss / 1048576))
            except (psutil.NoSuchProcess, psutil.AccessDenied):
                continue
    except Exception:
        pass
    return out


def kill_browser_processes(job_dir: Path):
    """Kill ONLY Chromium processes (never the Playwright node driver)."""
    if psutil is None:
        return
    killed = 0
    try:
        for p in psutil.Process(os.getpid()).children(recursive=True):
            try:
                ident = (p.name() + " " + " ".join(p.cmdline())).lower()
                if "chrom" in ident and not p.name().startswith("node"):
                    p.kill()
                    killed += 1
            except Exception:
                pass
        job_log(job_dir, f"Killed {killed} Chromium process(es).", "ERROR")
    except Exception as exc:
        job_log(job_dir, f"Could not kill browser processes: {exc}", "ERROR")


def collect_snapshot(dl_dir: Path) -> dict:
    cg = cgroup_memory()
    tree = process_tree()
    chrom = [t for t in tree if t[1].startswith("chromium")]
    used = cg["used_mb"]
    if used is None and tree:
        used = sum(t[3] for t in tree)
    ws = cg["ws_mb"] if cg["ws_mb"] is not None else used
    files = []
    try:
        files = [(p.name, p.stat().st_size / 1048576) for p in dl_dir.rglob("*") if p.is_file()]
    except OSError:
        pass
    try:
        disk_tmp = shutil.disk_usage(tempfile.gettempdir()).free / 1048576
    except OSError:
        disk_tmp = None
    try:
        disk_shm = shutil.disk_usage("/dev/shm").free / 1048576
    except OSError:
        disk_shm = None
    return {
        "used": used, "ws": ws, "limit": cg["limit_mb"], "peak": cg["peak_mb"], "oom": cg["oom_kills"],
        "anon": cg["anon_mb"], "file": cg["file_mb"], "shmem": cg["shmem_mb"],
        "py_rss": sum(t[3] for t in tree if t[1] == "streamlit"),
        "node_rss": sum(t[3] for t in tree if t[1] == "playwright-node"),
        "chrom_n": len(chrom), "chrom_rss": sum(t[3] for t in chrom),
        "chrom_top": sorted(chrom, key=lambda t: -t[3])[:4],
        "tree": tree, "files": files,
        "dl_mb": max((f[1] for f in files), default=0.0),
        "disk_tmp": disk_tmp, "disk_shm": disk_shm,
    }


def format_snapshot(s: dict) -> str:
    def f(v, unit="MB"):
        return "?" if v is None else f"{v:.0f}{unit}"

    pct = f" ({s['ws'] / s['limit']:.0%})" if s.get("ws") is not None and s["limit"] else ""
    top = ", ".join(f"{t[1].replace('chromium-', '')} {t[3]:.0f}" for t in s["chrom_top"]) or "none"
    return (
        f"WS={f(s.get('ws'))}/{f(s['limit'])}{pct} used={f(s['used'])} anon={f(s.get('anon'))} "
        f"file={f(s.get('file'))} shmem={f(s.get('shmem'))} peak={f(s['peak'])} oom_kills={s['oom']} | "
        f"streamlit={f(s.get('py_rss'))} node={f(s.get('node_rss'))} "
        f"chromium procs={s['chrom_n']} rss={f(s['chrom_rss'])} [{top}] | "
        f"free: tmp={f(s['disk_tmp'])} shm={f(s['disk_shm'])}"
    )


def log_system_info(job_dir: Path):
    try:
        from importlib.metadata import version

        pw_version = version("playwright")
    except Exception:
        pw_version = "?"
    cg = cgroup_memory()
    total = f"{psutil.virtual_memory().total / 1048576:.0f}MB" if psutil else "?"
    job_log(
        job_dir,
        f"SYSTEM: python={sys.version.split()[0]} playwright={pw_version} cpus={os.cpu_count()} "
        f"host_RAM={total} container_limit={cg['limit_mb'] and round(cg['limit_mb'])}MB "
        f"chromium={shutil.which('chromium') or shutil.which('chromium-browser')} psutil={'yes' if psutil else 'NO'}",
    )
    lim = cg["limit_mb"]
    if lim:
        job_log(job_dir, f"MEMORY POLICY: hard guard at {MEM_GUARD_FRACTION:.0%} of working set "
                         f"({lim * MEM_GUARD_FRACTION:.0f} MB, {MEM_GUARD_SAMPLES} samples in a row); "
                         f"soft recycle at {SOFT_RECYCLE_FRACTION:.0%} ({lim * SOFT_RECYCLE_FRACTION:.0f} MB) "
                         f"or Chromium +{CHROMIUM_GROWTH_MB} MB.")


class DiagnosticsMonitor(threading.Thread):
    """Samples memory / processes, logs them, and kills Chromium only under sustained real memory pressure."""

    def __init__(self, job_dir: Path, dl_dir: Path, rep: StepReporter):
        super().__init__(daemon=True)
        self.job_dir, self.dl_dir, self.rep = job_dir, dl_dir, rep
        self.stop_evt = threading.Event()
        self.t0 = time.time()
        self.phase = "starting"
        self.tripped = None
        self.peak_used = 0.0
        self.peak_chromium = 0.0
        self._seen_chromium = False
        self._gone_logged = False
        self._over = 0
        self._last_logged = 0.0
        self._last_used = 0.0

    def reset(self):
        """Called at the start of every browser session."""
        self.tripped = None
        self._over = 0
        self._seen_chromium = False
        self._gone_logged = False

    def run(self):
        while not self.stop_evt.wait(MONITOR_INTERVAL):
            try:
                self._tick()
            except Exception as exc:
                log.warning("monitor error: %s", exc)

    def stop(self):
        self.stop_evt.set()

    def _tick(self):
        now = time.time()
        elapsed = now - self.t0
        s = collect_snapshot(self.dl_dir)
        ws = s["ws"] or 0.0
        self.peak_used = max(self.peak_used, ws)
        self.peak_chromium = max(self.peak_chromium, s["chrom_rss"])

        if s["chrom_n"]:
            self._seen_chromium = True
        elif self._seen_chromium and not self._gone_logged and self.phase not in QUIET_PHASES:
            self._gone_logged = True
            job_log(self.job_dir, "!!! ALL CHROMIUM PROCESSES ARE GONE (browser exited or was killed). "
                                  f"Last memory state: {format_snapshot(s)}", "ERROR")

        if elapsed < 60 or now - self._last_logged >= 15 or ws - self._last_used >= 100:
            self._last_logged, self._last_used = now, ws
            job_log(self.job_dir, f"MON +{int(elapsed)}s [{self.phase}] {format_snapshot(s)}")

        limit = s["limit"]
        if self.phase in QUIET_PHASES:
            self._over = 0
        elif limit and ws > limit * MEM_GUARD_FRACTION:
            self._over += 1
        else:
            self._over = 0

        if limit and self._over >= MEM_GUARD_SAMPLES and not self.tripped:
            self.tripped = (
                f"Memory guard stopped Chromium: the container working set stayed above "
                f"{MEM_GUARD_FRACTION:.0%} of the limit ({ws:.0f} of {limit:.0f} MB) for "
                f"{MEM_GUARD_SAMPLES} samples. Chromium used {s['chrom_rss']:.0f} MB."
            )
            job_log(self.job_dir, f"!!! {self.tripped}", "ERROR")
            for pid, role, status, rss in sorted(s["tree"], key=lambda t: -t[3])[:8]:
                job_log(self.job_dir, f"    pid={pid} {role} status={status} rss={rss:.0f}MB", "ERROR")
            kill_browser_processes(self.job_dir)


def attach_diagnostics(job_dir: Path, browser, page):
    counters = {"console": 0, "net": 0}

    def guarded(fn):
        def wrapper(*args):
            try:
                fn(*args)
            except Exception:
                pass
        return wrapper

    @guarded
    def on_console(msg):
        if msg.type in ("error", "warning") and counters["console"] < 100:
            counters["console"] += 1
            job_log(job_dir, f"[console.{msg.type}] {msg.text[:300]}", "WARNING")

    @guarded
    def on_requestfailed(r):
        if counters["net"] < 30:
            counters["net"] += 1
            job_log(job_dir, f"[requestfailed] {r.method} {r.url[:150]} -> {r.failure}", "WARNING")

    browser.on("disconnected", guarded(lambda *_: job_log(
        job_dir, "BROWSER DISCONNECTED (closed or crashed)", "WARNING")))
    page.on("crash", guarded(lambda *_: job_log(
        job_dir, "!!! PAGE CRASHED: the renderer process died (usually out of memory)", "ERROR")))
    page.on("console", on_console)
    page.on("pageerror", guarded(lambda err: job_log(job_dir, f"[pageerror] {str(err)[:300]}", "WARNING")))
    page.on("requestfailed", on_requestfailed)


def take_shot(page, job_dir: Path, name: str):
    try:
        page.screenshot(path=str(job_dir / f"{name}.jpg"), type="jpeg", quality=50, timeout=15000)
        job_log(job_dir, f"Screenshot saved: {name}.jpg")
    except Exception as exc:
        job_log(job_dir, f"Screenshot {name} failed: {str(exc)[:150]}", "WARNING")


def _dispose(handle):
    try:
        handle.dispose()
    except Exception:
        pass


# =====================================================================
#  Playwright helpers
# =====================================================================
def launch_browser(pw, job_dir: Path, dl_dir: Path):
    args = [
        "--no-sandbox", "--disable-dev-shm-usage", "--disable-gpu", "--mute-audio",
        "--renderer-process-limit=1", "--no-zygote",
        "--disable-features=IsolateOrigins,site-per-process", "--disable-site-isolation-trials",
        "--disable-extensions", "--disable-background-networking", "--disable-sync",
        "--disable-component-update", "--disable-breakpad", "--metrics-recording-only", "--no-first-run",
    ]
    limit = cgroup_memory()["limit_mb"]
    if limit:
        args.append(f"--js-flags=--max-old-space-size={int(max(256, limit * 0.45))}")

    kwargs = dict(headless=True, downloads_path=str(dl_dir), args=args)

    system_chromium = shutil.which("chromium") or shutil.which("chromium-browser")
    if system_chromium:
        try:
            browser = pw.chromium.launch(executable_path=system_chromium, **kwargs)
            job_log(job_dir, f"Using system Chromium {system_chromium} (version {browser.version})")
            return browser
        except PWError as exc:
            job_log(job_dir, f"System Chromium failed ({str(exc)[:300]}); trying Playwright's own.", "WARNING")

    try:
        browser = pw.chromium.launch(**kwargs)
        job_log(job_dir, f"Using Playwright Chromium (version {browser.version})")
        return browser
    except PWError as exc:
        if "playwright install" not in str(exc) and "Executable doesn't exist" not in str(exc):
            raise
    with _INSTALL_LOCK:
        job_log(job_dir, "Installing Playwright Chromium (first run only, may take a minute)...")
        proc = subprocess.run(
            [sys.executable, "-m", "playwright", "install", "chromium"],
            capture_output=True, text=True, timeout=900,
        )
        if proc.returncode != 0:
            raise RuntimeError(f"'playwright install chromium' failed: {proc.stderr[-500:] or proc.stdout[-500:]}")
    return pw.chromium.launch(**kwargs)


def first_visible(page, selector, timeout=DEFAULT_TIMEOUT):
    """Return the first VISIBLE match of a selector (the hidden 'Active Alarm' tab may contain look-alikes)."""
    loc = page.locator(selector)
    end = time.time() + timeout
    while True:
        try:
            n = loc.count()
            for i in range(n):
                cand = loc.nth(i)
                if cand.is_visible():
                    return cand
        except PWError:
            pass
        if time.time() > end:
            raise RuntimeError(f"No visible element found for: {selector}")
        page.wait_for_timeout(300)


def safe_click(loc, timeout=10000):
    try:
        loc.scroll_into_view_if_needed(timeout=timeout)
    except PWError:
        pass
    try:
        loc.click(timeout=timeout)
    except PWError:
        loc.evaluate("el => el.click()")


def click_first(page, selectors, description, total_timeout=DEFAULT_TIMEOUT):
    per = max(5, int(total_timeout / max(1, len(selectors))))
    last_exc = None
    for sel in selectors:
        try:
            safe_click(first_visible(page, sel, per))
            return
        except (PWError, RuntimeError) as exc:
            last_exc = exc
            log.warning("Selector failed for %s: %s", description, sel)
    raise RuntimeError(f"Could not click {description}: {str(last_exc)[:200]}")


UPPER = "translate(normalize-space(.), 'abcdefghijklmnopqrstuvwxyz', 'ABCDEFGHIJKLMNOPQRSTUVWXYZ')"


def pw_login(page, username, password):
    page.goto(LOGIN_URL, wait_until="load", timeout=60000)
    user = page.locator(
        "input[placeholder='Username'], input[name='username'], input#username"
    ).first
    user.wait_for(state="visible")
    user.fill(username)
    pwd = page.locator("input[type='password']").first
    pwd.wait_for(state="visible")
    pwd.fill(password)

    click_first(
        page,
        [
            "xpath=//button[normalize-space()='LOGIN' or normalize-space()='Login']",
            "xpath=//button[contains(translate(., 'login', 'LOGIN'), 'LOGIN')]",
            "xpath=//input[@type='submit']",
            "css=button[type='submit']",
        ],
        "LOGIN button",
    )
    try:
        page.wait_for_url(lambda url: "/auth/login" not in url, timeout=60000)
    except PWTimeout:
        raise RuntimeError(
            "Login failed: still on the login page. Check username/password, "
            "or the site may require CAPTCHA/OTP or block this server's IP."
        )


def pw_open_alarm_page(page):
    page.goto(ALARM_URL, wait_until="load", timeout=60000)
    if "/auth/login" in page.url:
        raise RuntimeError("Redirected back to login page; session not authenticated.")
    page.wait_for_selector("body", state="attached")
    page.wait_for_timeout(2000)


CLOSED_TAB_XPATHS = [
    "xpath=//*[(self::button or self::a or self::li or self::div or self::span or @role='tab')]"
    "[normalize-space()='CLOSED ALARM' or normalize-space()='Closed Alarm']",
    f"xpath=//*[contains({UPPER}, 'CLOSED ALARM') and (self::button or self::a or @role='tab' or self::li)]",
]


def _closed_tab_visible(page) -> bool:
    for sel in CLOSED_TAB_XPATHS:
        try:
            loc = page.locator(sel)
            for i in range(loc.count()):
                if loc.nth(i).is_visible():
                    return True
        except PWError:
            pass
    return False


def pw_scroll_to_alarm_section(page):
    """Scroll down the dashboard (past Daily Analytics) until the ACTIVE / CLOSED ALARM tabs are on screen."""
    section = page.locator(
        "xpath=//*[self::h1 or self::h2 or self::h3 or self::h4 or self::h5 or self::div or self::span]"
        "[normalize-space()='Alarm' or normalize-space()='ALARM']"
    ).first
    try:
        section.wait_for(state="attached", timeout=5000)
        section.evaluate("el => el.scrollIntoView({block: 'start'})")
    except PWError:
        pass
    for _ in range(25):  # keep scrolling until the tab appears
        if _closed_tab_visible(page):
            return
        page.mouse.wheel(0, 700)
        page.wait_for_timeout(500)
    page.evaluate("window.scrollTo(0, document.body.scrollHeight)")
    page.wait_for_timeout(800)


def wait_loading_done(page, timeout=10000):
    try:
        page.wait_for_selector(
            ".spinner, .loading, .loader, .spinner-border, [class*='loading'], [class*='spinner']",
            state="hidden", timeout=timeout,
        )
    except PWError:
        pass


def pw_click_closed_alarm_tab(page, job_dir: Path):
    click_first(page, CLOSED_TAB_XPATHS, "CLOSED ALARM tab")
    # wait until the Closed Alarm table (header has Site ID + Severity) is really on screen
    end = time.time() + 60
    headers = []
    while time.time() < end:
        data = page.evaluate(READ_TABLE_JS)
        if data and data["heads"]:
            headers = data["heads"]
            if data["rows"]:
                break
        page.wait_for_timeout(500)
    if not headers:
        raise RuntimeError("The CLOSED ALARM table did not appear after clicking the tab.")
    job_log(job_dir, f"Closed Alarm table detected, columns: {headers}")
    wait_loading_done(page, 20000)
    page.wait_for_timeout(4000)   # the pagination footer may render later than the table


# =====================================================================
#  Filter controls: Start Date / Apply / Items per page   (End Date is never touched)
# =====================================================================
# The page also has a "Daily Analytics" section with its own Start/End Date boxes and tables.
# Everything below is anchored to the CLOSED ALARM table (header contains Site ID + Severity),
# never to "the first visible thing with that label".
_FIND_TABLE = r"""
const findTable = () => {
  const tables = Array.from(document.querySelectorAll('table')).filter(t => t.offsetParent !== null);
  const head = t => ((t.tHead ? t.tHead.innerText : '') || '').toLowerCase();
  return tables.find(t => /site\s*id/.test(head(t)) && /severity/.test(head(t)))
      || tables.find(t => /site\s*id/.test(head(t)))
      || null;
};
"""

MARK_CONTROL_JS = "(kind) => {" + _FIND_TABLE + r"""
  document.querySelectorAll('[data-pw-target]').forEach(e => e.removeAttribute('data-pw-target'));
  const table = findTable();
  if (!table) return 'no-closed-alarm-table';
  const vis = el => el.offsetParent !== null || el.getClientRects().length > 0;
  const own = el => Array.from(el.childNodes).filter(n => n.nodeType === 3)
                         .map(n => n.textContent).join(' ').replace(/[*:]/g, '').replace(/\s+/g, ' ').trim().toLowerCase();
  const before = el => !table.contains(el) && (el.compareDocumentPosition(table) & Node.DOCUMENT_POSITION_FOLLOWING);
  const after = el => !table.contains(el) && (table.compareDocumentPosition(el) & Node.DOCUMENT_POSITION_FOLLOWING);
  const all = Array.from(document.querySelectorAll('body *'));
  const tableTop = table.getBoundingClientRect().top;
  let target = null;
  if (kind === 'start' || kind === 'end') {
    const text = kind === 'start' ? 'start date' : 'end date';
    // the label that sits closest ABOVE the Closed Alarm table
    const labels = all.filter(el => vis(el) && before(el) && own(el) === text);
    const label = labels[labels.length - 1];
    if (!label) return 'no-label';
    const inputs = Array.from(document.querySelectorAll('input'))
      .filter(i => vis(i) && (label.contains(i) || (label.compareDocumentPosition(i) & Node.DOCUMENT_POSITION_FOLLOWING)));
    target = inputs[0];
  } else if (kind === 'apply') {
    const btns = Array.from(document.querySelectorAll('button, [role="button"]'))
      .filter(el => vis(el) && before(el) && (el.innerText || '').replace(/\s+/g, ' ').trim().toUpperCase() === 'APPLY');
    target = btns[btns.length - 1];
  } else if (kind === 'items') {
    const labels = all.filter(el => vis(el) && after(el) && own(el) === 'items');
    const label = labels[0];
    if (!label) return 'no-label';
    const cands = Array.from(document.querySelectorAll('select, [role="combobox"], mat-select, [class*="select"]'))
      .filter(el => vis(el) && (label.compareDocumentPosition(el) & Node.DOCUMENT_POSITION_FOLLOWING));
    target = cands[0];
  }
  if (!target) return 'no-target';
  target.setAttribute('data-pw-target', kind);
  const d = Math.round(Math.abs(target.getBoundingClientRect().top - tableTop));
  return 'ok:' + d;
}"""


def mark_control(page, kind: str, timeout=DEFAULT_TIMEOUT):
    """Locate a filter control that belongs to the CLOSED ALARM table (not Daily Analytics)."""
    end = time.time() + timeout
    res = ""
    while time.time() < end:
        res = page.evaluate(MARK_CONTROL_JS, kind)
        if str(res).startswith("ok"):
            return page.locator(f"[data-pw-target='{kind}']").first
        page.wait_for_timeout(400)
    raise RuntimeError(f"Could not find the '{kind}' control next to the Closed Alarm table (result: {res}).")


def pw_click_apply(page):
    safe_click(mark_control(page, "apply"))


READ_TABLE_JS = "() => {" + _FIND_TABLE + r"""
  const clean = s => (s || '').replace(/\s+/g, ' ').trim();
  const table = findTable();
  if (!table) return null;
  const heads = Array.from(table.querySelectorAll('thead th')).map(th => clean(th.innerText));
  const rows = Array.from(table.querySelectorAll('tbody tr')).map(tr =>
      Array.from(tr.querySelectorAll('td')).map(td => clean(td.innerText)));
  return {heads, rows};
}"""

COUNT_ROWS_JS = "() => {" + _FIND_TABLE + r"""
  const table = findTable();
  if (!table) return -1;
  return Array.from(table.querySelectorAll('tbody tr'))
      .filter(tr => tr.querySelectorAll('td').length >= 3 && (tr.innerText || '').trim()).length;
}"""

# Tiny signature of the table (row count + first row + last row). Computed INSIDE the browser,
# so a 500-row table is never copied to Python while we wait for a page change.
SIG_JS = "() => {" + _FIND_TABLE + r"""
  const t = findTable();
  if (!t) return '';
  const tr = t.querySelectorAll('tbody tr');
  if (!tr.length) return '';
  const c = e => (e.innerText || '').replace(/\s+/g, ' ').trim();
  return tr.length + '|' + c(tr[0]) + '|' + c(tr[tr.length - 1]);
}"""

TOTAL_PAGES_JS = "() => {" + _FIND_TABLE + r"""
  const table = findTable();
  if (!table) return 0;
  let max = 0;
  document.querySelectorAll('button, a, li, [role="button"]').forEach(el => {
    if (table.contains(el) || !(table.compareDocumentPosition(el) & Node.DOCUMENT_POSITION_FOLLOWING)) return;
    if (el.closest('[role="listbox"], [role="option"], [role="menu"], [role="combobox"], select, .v-select, .v-field, ' +
                   '.v-overlay, .v-menu, .MuiPopover-root, .MuiSelect-select, .ant-select, mat-select, ' +
                   '.cdk-overlay-container')) return;
    const t = (el.innerText || '').trim();
    if (/^\d{1,5}$/.test(t)) max = Math.max(max, parseInt(t, 10));
  });
  return max;
}"""

FIND_NEXT_JS = "() => {" + _FIND_TABLE + r"""
  const table = findTable();
  if (!table) return null;
  const after = el => !table.contains(el) && (table.compareDocumentPosition(el) & Node.DOCUMENT_POSITION_FOLLOWING);
  const vis = el => el.offsetParent !== null || el.getClientRects().length > 0;
  const cands = Array.from(document.querySelectorAll('button, a, [role="button"]')).filter(el => after(el) && vis(el));
  if (!cands.length) return null;
  const txt = el => (el.innerText || '').trim();
  const cls = el => (el.className && el.className.toString) ? el.className.toString() : '';
  const meta = el => ((el.getAttribute('aria-label') || '') + ' ' + (el.getAttribute('title') || '') + ' ' + cls(el)).toLowerCase();
  const byMeta = cands.filter(el => /next|chevron_right|chevron-right|arrow_forward/.test(meta(el)) && !/last|prev|first/.test(meta(el)));
  if (byMeta.length) return byMeta[byMeta.length - 1];
  const byText = cands.filter(el => /^(chevron_right|navigate_next|keyboard_arrow_right|arrow_forward_ios|arrow_right|›|>|»)$/i.test(txt(el)));
  if (byText.length) return byText[byText.length - 1];
  const noDigit = cands.filter(el => !/\d/.test(txt(el)));
  return noDigit.length ? noDigit[noDigit.length - 1] : null;
}"""

IS_DISABLED_JS = r"""
el => {
  const cls = n => ((n && n.className && n.className.toString) ? n.className.toString() : '').toLowerCase();
  const p = el.parentElement;
  return !!(el.disabled || el.hasAttribute('disabled') || el.getAttribute('aria-disabled') === 'true'
            || /disabled/.test(cls(el)) || (p && /disabled/.test(cls(p))));
}
"""

# Page-number buttons below the Closed Alarm table. Marks the one to click with data-pw-page.
#   mode 'exact' : the button with the wanted number is visible
#   mode 'jump'  : it is not visible, so the visible button closest BELOW the target (and above the current page)
#   mode 'none'  : nothing useful -> the caller clicks "Next"
# `cur` is the active page number (0 = could not be detected). Call with target=0 to only read `cur`.
MARK_PAGE_JS = "(target) => {" + _FIND_TABLE + r"""
  document.querySelectorAll('[data-pw-page]').forEach(e => e.removeAttribute('data-pw-page'));
  const table = findTable();
  if (!table) return {mode: 'no-table', cur: 0, digits: []};
  const after = el => !table.contains(el) && (table.compareDocumentPosition(el) & Node.DOCUMENT_POSITION_FOLLOWING);
  const vis = el => el.offsetParent !== null || el.getClientRects().length > 0;
  const tokens = e => ((e && e.className && e.className.toString) ? e.className.toString() : '').split(/\s+/);
  const isAct = e => {
    if (!e || !e.getAttribute) return false;
    const ac = e.getAttribute('aria-current');
    if (ac && ac !== 'false') return true;
    if (e.getAttribute('aria-selected') === 'true') return true;
    return tokens(e).some(t => /(^|[-_])(active|selected|current)$/i.test(t));
  };
  const EXCL = '[role="listbox"], [role="option"], [role="menu"], [role="combobox"], select, .v-select, .v-field, ' +
               '.v-overlay, .v-menu, .MuiPopover-root, .MuiSelect-select, .ant-select, mat-select, .cdk-overlay-container';
  let els = Array.from(document.querySelectorAll('button, a, li, [role="button"]'))
      .filter(el => after(el) && vis(el) && !el.closest(EXCL) && /^\d{1,5}$/.test((el.innerText || '').trim()));
  els = els.filter(el => !els.some(o => o !== el && el.contains(o)));   // keep the deepest element only
  const digits = els.map(el => ({n: parseInt(el.innerText.trim(), 10), el: el,
                                 act: isAct(el) || isAct(el.parentElement)}));
  const actEl = digits.find(d => d.act);
  const cur = actEl ? actEl.n : 0;
  let pick = digits.find(d => d.n === target), mode = 'exact';
  if (!pick) {
    pick = digits.filter(d => d.n < target && d.n > cur).sort((a, b) => b.n - a.n)[0];
    mode = 'jump';
  }
  if (!pick) return {mode: 'none', cur: cur, digits: digits.map(d => d.n)};
  pick.el.setAttribute('data-pw-page', '1');
  return {mode: mode, cur: cur, digits: digits.map(d => d.n)};
}"""

_ICON_TAIL_RE = re.compile(r"\s+(info|info_outline|help|help_outline|arrow_upward|arrow_downward|unfold_more)$", re.I)


def read_table(page):
    """Return (headers, rows) of the visible table. Only used for the initial checks (it copies the whole table)."""
    data = page.evaluate(READ_TABLE_JS)
    if not data:
        return [], []
    rows = [r for r in data["rows"] if len(r) >= 3 and any(r)]   # skip 'no data' placeholder rows
    width = max([len(data["heads"])] + [len(r) for r in rows])
    seen, headers = {}, []
    for i in range(width):
        h = _ICON_TAIL_RE.sub("", data["heads"][i]).strip() if i < len(data["heads"]) else ""
        h = h or f"col_{i + 1}"
        n = seen.get(h, 0)
        seen[h] = n + 1
        headers.append(h if n == 0 else f"{h}_{n + 1}")
    return headers, rows


def count_rows(page) -> int:
    """Cheap row counter (does not copy every cell back to Python, so it is fine for 500-row tables)."""
    try:
        return max(0, int(page.evaluate(COUNT_ROWS_JS)))
    except (PWError, TypeError, ValueError):
        return 0


def get_sig(page) -> str:
    try:
        return page.evaluate(SIG_JS) or ""
    except PWError:
        return ""


def wait_sig_change(page, old, timeout=40, poll_ms=700):
    """Light-weight change detection: waits until the table signature differs from `old` and is stable."""
    end = time.time() + timeout
    while time.time() < end:
        s = get_sig(page)
        if s and s != old:
            page.wait_for_timeout(500)
            if get_sig(page) == s:
                return True
            continue
        page.wait_for_timeout(poll_ms)
    return False


def wait_table_stable(page, timeout=30, still=STABLE_SECONDS) -> bool:
    """True once the table signature stayed identical for `still` seconds."""
    end = time.time() + timeout
    last, since = None, time.time()
    while time.time() < end:
        s = get_sig(page)
        if s != last:
            last, since = s, time.time()
        elif time.time() - since >= still:
            return True
        page.wait_for_timeout(500)
    return False


# =====================================================================
#  REAL CALENDAR (date-picker) for the Start Date
#  Opens the widget by clicking the field, navigates month/year with the widget's own buttons, clicks the day.
#  Supports Vuetify, flatpickr, MUI, Ant Design, bootstrap-datepicker, react-datepicker, Quasar, Angular
#  Material and generic widgets. If it cannot read the widget it dumps the HTML to job.log + dom_cal_diag.txt.
# =====================================================================
_CAL_COMMON = r"""
const CAL_CSS = ".v-date-picker, .v-picker, .v-date-picker-month, .flatpickr-calendar, .MuiDateCalendar-root, " +
  ".MuiPickersPopper-root, .MuiPickerStaticWrapper-root, .ant-picker-dropdown, .ant-picker-panel, " +
  ".datepicker, .datepicker-dropdown, .react-datepicker, .react-datepicker-popper, .q-date, .mat-calendar, " +
  ".mat-datepicker-content, .dp__menu, .vc-container, .daterangepicker, [class*='datepicker' i], " +
  "[class*='date-picker' i], [class*='calendar' i], [role='dialog']";
const MON = {jan:0,feb:1,mar:2,apr:3,mei:4,may:4,jun:5,jul:6,agu:7,agt:7,ags:7,aug:7,sep:8,okt:9,oct:9,nov:10,des:11,dec:11};
const MON_RE = /\b(jan|feb|mar|apr|mei|may|jun|jul|agu|agt|ags|aug|sep|okt|oct|nov|des|dec)[a-z]*\.?,?\s*(20\d\d)\b/i;
const vis = el => (el.offsetParent !== null || el.getClientRects().length > 0) && el.getBoundingClientRect().width > 0;
const dayCount = e => Array.from(e.querySelectorAll('button, td, [role="gridcell"], span, div'))
    .filter(x => x.children.length <= 1 && /^\d{1,2}$/.test((x.innerText || '').trim())).length;
const findCal = () => {
  const c = Array.from(document.querySelectorAll(CAL_CSS)).filter(e => {
    if (!vis(e)) return false;
    const r = e.getBoundingClientRect();
    return r.width > 150 && r.width < 700 && r.height > 150 && r.height < 700 && dayCount(e) >= 28;
  });
  c.sort((a, b) => b.getBoundingClientRect().width * b.getBoundingClientRect().height -
                   a.getBoundingClientRect().width * a.getBoundingClientRect().height);
  return c[0] || null;      // the largest matching container = includes the header and the prev/next buttons
};
const clear = () => document.querySelectorAll('[data-pw-cal]').forEach(e => e.removeAttribute('data-pw-cal'));
"""

CAL_STATE_JS = "() => {" + _CAL_COMMON + r"""
  const c = findCal();
  if (!c) return null;
  let month = -1, year = 0;
  // flatpickr style: <select> with months + <input> with the year
  const sel = Array.from(c.querySelectorAll('select')).find(s =>
      Array.from(s.options).some(o => MON[(o.text || '').trim().slice(0, 3).toLowerCase()] !== undefined));
  const yin = Array.from(c.querySelectorAll('input[type="number"], input.cur-year')).find(i => /^20\d\d$/.test(i.value));
  if (sel && yin) {
    month = MON[(sel.options[sel.selectedIndex].text || '').trim().slice(0, 3).toLowerCase()];
    year = parseInt(yin.value, 10);
  } else {
    const m = (c.innerText || '').match(MON_RE);
    if (m) { month = MON[m[1].toLowerCase().slice(0, 3)]; year = parseInt(m[2], 10); }
  }
  return {month: month === undefined ? -1 : month, year: year, head: (c.innerText || '').slice(0, 80)};
}"""

CAL_MARK_NAV_JS = "(dir) => {" + _CAL_COMMON + r"""
  clear();
  const c = findCal();
  if (!c) return false;
  const meta = e => ((e.getAttribute('aria-label') || '') + ' ' + (e.getAttribute('title') || '') + ' ' +
                     ((e.className && e.className.toString) ? e.className.toString() : '') + ' ' +
                     (e.innerText || '')).toLowerCase();
  const re = dir === 'prev' ? /prev|chevron[-_]?left|arrow[-_]?left|navigate_before|keyboard_arrow_left|‹|«|^\s*<\s*$/
                            : /next|chevron[-_]?right|arrow[-_]?right|navigate_next|keyboard_arrow_right|›|»|^\s*>\s*$/;
  const bad = e => e.disabled || e.getAttribute('aria-disabled') === 'true' || /disabled/.test(meta(e));
  const cands = Array.from(c.querySelectorAll('button, [role="button"], span, i, a, th, div'))
      .filter(e => vis(e) && re.test(meta(e)) && !/^\d{1,2}$/.test((e.innerText || '').trim()) && !bad(e));
  const leaf = cands.filter(e => !cands.some(o => o !== e && e.contains(o)));   // deepest only
  const btn = leaf.find(e => e.closest('button, [role="button"]')) || leaf[0];
  if (!btn) return false;
  (btn.closest('button, [role="button"]') || btn).setAttribute('data-pw-cal', 'nav');
  return true;
}"""

CAL_MARK_DAY_JS = "(day) => {" + _CAL_COMMON + r"""
  clear();
  const c = findCal();
  if (!c) return 'no-calendar';
  const OUT = /(^|[-_ ])(adjacent|outside|outside-month|othermonth|other-month|prevmonthday|nextmonthday|prevmonth|nextmonth|old|new|disabled|muted)([-_ ]|$)/i;
  const cls = e => ((e.className && e.className.toString) ? e.className.toString() : '');
  const excluded = e => {
    let n = e;
    for (let i = 0; i < 4 && n && n !== c; i++, n = n.parentElement) {
      if (OUT.test(cls(n)) || n.disabled || n.getAttribute('aria-disabled') === 'true') return true;
      if (/ant-picker-cell/.test(cls(n)) && !/in-view/.test(cls(n))) return true;
    }
    return false;
  };
  const all = Array.from(c.querySelectorAll('button, td, [role="gridcell"], span, div, a'))
      .filter(e => vis(e) && (e.innerText || '').trim() === String(day) && !excluded(e));
  const leaf = all.filter(e => !all.some(o => o !== e && e.contains(o)));
  const pick = leaf.find(e => e.closest('button')) || leaf[0];
  if (!pick) return 'no-day-cell (' + all.length + ' candidates)';
  pick.setAttribute('data-pw-cal', 'day');
  return 'ok';
}"""

CAL_MARK_OK_JS = "() => {" + _CAL_COMMON + r"""
  clear();
  const c = findCal();
  if (!c) return false;
  const b = Array.from(c.querySelectorAll('button, [role="button"]')).filter(e => vis(e) &&
      /^(ok|set|select|done|choose|pilih|simpan|confirm)$/i.test((e.innerText || '').trim()));
  if (!b.length) return false;
  b[b.length - 1].setAttribute('data-pw-cal', 'ok');
  return true;
}"""

CAL_HTML_JS = "() => {" + _CAL_COMMON + r"""
  const c = findCal();
  return c ? c.outerHTML.slice(0, 3500) : '(no calendar container found)';
}"""

CAL_OPEN_JS = "() => {" + _CAL_COMMON + "  return !!findCal();\n}"


def _click_real(page, name: str):
    """Genuine mouse click on the element marked data-pw-cal=<name> (no JS click, no value injection)."""
    loc = page.locator(f"[data-pw-cal='{name}']").first
    try:
        loc.scroll_into_view_if_needed(timeout=3000)
        loc.click(timeout=6000)
    except PWError:
        box = loc.bounding_box()
        if not box:
            raise RuntimeError(f"calendar element '{name}' has no bounding box")
        page.mouse.click(box["x"] + box["width"] / 2, box["y"] + box["height"] / 2)


def _calendar_open(page) -> bool:
    try:
        return bool(page.evaluate(CAL_OPEN_JS))
    except PWError:
        return False


def append_dom(job_dir: Path, tag: str, text: str):
    try:
        with open(job_dir / f"dom_{tag}diag.txt", "a", encoding="utf-8") as fh:
            fh.write(text + "\n")
    except OSError:
        pass


def _dump_calendar(page, job_dir: Path, why: str):
    try:
        html = page.evaluate(CAL_HTML_JS)
    except PWError:
        html = "(could not read)"
    job_log(job_dir, f"CALENDAR DUMP ({why}): {html[:1500]}", "ERROR")
    append_dom(job_dir, "cal_", f"CALENDAR DUMP ({why}):\n{html}")
    take_shot(page, job_dir, "shot_calendar_problem")


def open_calendar(page, job_dir: Path):
    """Click the Start Date field like a user; if no calendar opens, click the calendar icon next to it."""
    inp = mark_control(page, "start", 15)
    inp.scroll_into_view_if_needed()
    inp.click()                                           # real click on the field
    for _ in range(10):
        if _calendar_open(page):
            return
        page.wait_for_timeout(300)
    # fallback: the icon/button inside the same field wrapper
    for xp in ("xpath=ancestor::*[self::div or self::label][1]//*[self::button or self::i or self::svg or "
               "contains(@class,'icon')]",
               "xpath=ancestor::*[self::div or self::label][2]//*[self::button or self::i or self::svg or "
               "contains(@class,'icon')]"):
        try:
            icon = mark_control(page, "start", 5).locator(xp).first
            if icon.count() and icon.is_visible():
                icon.click(timeout=4000)
                for _ in range(10):
                    if _calendar_open(page):
                        job_log(job_dir, "[Start Date] calendar opened via the calendar icon.")
                        return
                    page.wait_for_timeout(300)
        except PWError:
            continue
    _dump_calendar(page, job_dir, "calendar did not open")
    raise RuntimeError("The Start Date calendar did not open (see CALENDAR DUMP in job.log and "
                       "shot_calendar_problem.jpg).")


def pw_pick_start_date(page, job_dir: Path, value: date):
    """
    Select the Start Date through the REAL calendar: open it, navigate months, click the day, (confirm),
    verify the field. No fill / typing / keyboard entry / JS value injection.
    """
    text = value.strftime(DATE_INPUT_FORMAT)
    ok_values = {re.sub(r"\D", "", text), value.strftime("%Y%m%d")}
    field = mark_control(page, "start", 15)
    job_log(job_dir, f"[Start Date] field shows '{field.input_value()}' before; selecting {text} in the calendar")

    open_calendar(page, job_dir)

    # ---- navigate to the wanted month/year
    target = value.year * 12 + (value.month - 1)
    for step in range(CAL_MAX_MONTH_CLICKS + 1):
        st_ = page.evaluate(CAL_STATE_JS)
        if not st_ or st_["month"] < 0 or not st_["year"]:
            _dump_calendar(page, job_dir, "cannot read month/year")
            raise RuntimeError(f"Cannot read the month/year shown by the calendar (header: "
                               f"{st_ and st_.get('head')!r}).")
        cur = st_["year"] * 12 + st_["month"]
        if cur == target:
            job_log(job_dir, f"[Start Date] calendar shows month {st_['month'] + 1}/{st_['year']} after "
                             f"{step} click(s).")
            break
        if step == CAL_MAX_MONTH_CLICKS:
            raise RuntimeError("Calendar navigation did not reach the wanted month.")
        direction = "prev" if cur > target else "next"
        if not page.evaluate(CAL_MARK_NAV_JS, direction):
            _dump_calendar(page, job_dir, f"no '{direction}' button")
            raise RuntimeError(f"Calendar has no usable '{direction}' month button.")
        _click_real(page, "nav")
        page.wait_for_timeout(250)
        st2 = page.evaluate(CAL_STATE_JS)
        if st2 and (st2["month"], st2["year"]) == (st_["month"], st_["year"]):
            page.wait_for_timeout(500)                      # animation may still be running

    # ---- click the day
    res = page.evaluate(CAL_MARK_DAY_JS, value.day)
    if res != "ok":
        _dump_calendar(page, job_dir, f"day {value.day}: {res}")
        raise RuntimeError(f"Could not find day {value.day} in the calendar ({res}).")
    _click_real(page, "day")
    page.wait_for_timeout(500)

    # ---- some pickers need an OK / Set button
    if _calendar_open(page) and page.evaluate(CAL_MARK_OK_JS):
        _click_real(page, "ok")
        page.wait_for_timeout(500)

    # ---- verify what the field really holds
    got = ""
    for _ in range(10):
        got = mark_control(page, "start", 10).input_value()
        if re.sub(r"\D", "", got) in ok_values:
            break
        page.wait_for_timeout(300)
    else:
        _dump_calendar(page, job_dir, "wrong value after click")
        raise RuntimeError(f"After clicking {text} in the calendar the field shows '{got}'.")

    if _calendar_open(page):                                 # popup still open: close it (value is kept)
        page.keyboard.press("Escape")
        page.wait_for_timeout(400)
        got = mark_control(page, "start", 10).input_value()
        if re.sub(r"\D", "", got) not in ok_values:
            raise RuntimeError(f"Closing the calendar changed the Start Date to '{got}'.")
    job_log(job_dir, f"[Start Date] SELECTED via calendar: '{got}' (wanted {text})")


def read_end_date(page) -> str:
    """READ-ONLY look at the End Date box (no click, no typing) for the log."""
    try:
        return mark_control(page, "end", 5).input_value()
    except (PWError, RuntimeError):
        return "?"


def ensure_dates(page, job_dir: Path, start_d: date, end_d: date):
    """
    Start Date: always selected through the calendar widget.
    End Date  : NEVER touched. The site keeps its own default (today). `end_d` is accepted only so that the
                callers do not change; it is not used here.
    """
    pw_pick_start_date(page, job_dir, start_d)
    job_log(job_dir, f"[End Date] left untouched, site value: '{read_end_date(page)}'")


# ---------------------------------------------------------------------
#  "Items" dropdown (10 / 50 / 100 / 250 / 500) -- FOOTER DISCOVERY EDITION
#  (locator logic unchanged: it already works; only the logging can be silenced for the per-page cycles)
# ---------------------------------------------------------------------
SIZE_ORDER = ["10", "50", "100", "250", "500"]
SITE_INFO = {"framework": "generic", "virtualized": False, "quiet": False}


def llog(job_dir: Path, msg: str, level: str = "INFO"):
    """job_log that can be silenced (INFO only) during the 250+ per-page cycles."""
    if SITE_INFO.get("quiet") and level == "INFO":
        return
    job_log(job_dir, msg, level)


SELECT_LIKE_CSS = (
    "select, [role='combobox'], [role='listbox'], [aria-haspopup='listbox'], mat-select, "
    ".v-select, .v-field, .MuiSelect-select, .MuiTablePagination-select, .ant-select, .q-select, "
    "[class*='select' i], [class*='dropdown' i], [class*='combo' i]"
)

FRAMEWORK_SELECTORS = {
    "vuetify_app": ".v-application, .v-app",
    "vuetify_table": ".v-data-table",
    "vuetify2_footer": ".v-data-footer",
    "vuetify3_footer": ".v-data-table-footer",
    "vuetify_select": ".v-select",
    "vuetify_virtual": ".v-data-table-virtual, .v-virtual-scroll",
    "mui_table": ".MuiTable-root",
    "mui_pagination": ".MuiTablePagination-root",
    "mui_datagrid": ".MuiDataGrid-root",
    "mui_virtual": ".MuiDataGrid-virtualScroller",
    "mat_paginator": "mat-paginator, .mat-mdc-paginator",
    "ant_table": ".ant-table",
    "ant_pagination": ".ant-pagination",
    "quasar": ".q-table",
    "ag_grid": ".ag-root",
    "virtual_strict": ".vue-recycle-scroller, .vue-virtual-scroller, [class*='virtual-scroll' i], "
                      "[class*='react-window' i], [class*='virtuoso' i]",
    "virtual_loose": "[class*='virtual' i]",
    "plain_table": "table",
    "tfoot": "tfoot",
}

# (name, css) -> tried in this order with Playwright locators. Framework-specific first, then generic.
ITEMS_SELECTORS = [
    ("vuetify3 items-per-page", ".v-data-table-footer__items-per-page .v-select, "
                                ".v-data-table-footer__items-per-page .v-field"),
    ("vuetify2 footer select", ".v-data-footer__select .v-select"),
    ("mui table pagination", ".MuiTablePagination-select, .MuiTablePagination-root [role='combobox']"),
    ("angular-material paginator", "mat-paginator mat-select, .mat-mdc-paginator-page-size-select mat-select, "
                                   ".mat-paginator-page-size-select mat-select"),
    ("antd pagination options", ".ant-pagination-options .ant-select"),
    ("quasar table bottom", ".q-table__bottom .q-select"),
    ("aria-label per page", "[aria-label*='per page' i], [aria-label*='items' i][role='combobox'], "
                            "[aria-label*='rows' i][role='combobox']"),
]

FRAMEWORK_JS = """(sels) => {
  const out = {};
  for (const [k, s] of Object.entries(sels)) {
    try { out[k] = document.querySelectorAll(s).length; } catch (e) { out[k] = 'err'; }
  }
  return out;
}"""

SCROLL_FOOTER_JS = "() => {" + _FIND_TABLE + r"""
  const out = {containers: [], win: null, noTable: false};
  const table = findTable();
  const scrollable = e => {
    const s = getComputedStyle(e);
    return /(auto|scroll|overlay)/.test(s.overflowY) && e.scrollHeight > e.clientHeight + 2;
  };
  const describe = e => e.tagName.toLowerCase() + (e.id ? '#' + e.id : '') +
      ((e.className && e.className.toString && e.className.toString().trim())
        ? '.' + e.className.toString().trim().split(/\s+/).slice(0, 3).join('.') : '');
  if (table) {
    const rows = table.querySelectorAll('tbody tr');
    if (rows.length) rows[rows.length - 1].scrollIntoView({block: 'end', behavior: 'instant'});
    let p = table.parentElement;
    while (p && p !== document.body && p !== document.documentElement) {
      if (scrollable(p)) {
        const before = Math.round(p.scrollTop);
        p.scrollTop = p.scrollHeight;
        out.containers.push({el: describe(p), before: before, after: Math.round(p.scrollTop),
                             scrollH: p.scrollHeight, clientH: p.clientHeight});
      }
      p = p.parentElement;
    }
  } else {
    out.noTable = true;
  }
  const se = document.scrollingElement || document.documentElement;
  const b = Math.round(se.scrollTop);
  se.scrollTop = se.scrollHeight;
  window.scrollTo(0, document.body.scrollHeight);
  out.win = {before: b, after: Math.round(se.scrollTop), scrollH: se.scrollHeight, innerH: window.innerHeight};
  return out;
}"""

FOOTER_DUMP_JS = "() => {" + _FIND_TABLE + r"""
  const table = findTable();
  if (!table) return {error: 'no Closed Alarm table found'};
  const cls = e => (e.className && e.className.toString) ? e.className.toString().trim() : '';
  const out = {chain: [], tfoot: table.tFoot ? table.tFoot.outerHTML.slice(0, 6000) : '', following: []};
  let child = table, parent = table.parentElement, level = 0;
  while (parent && level < 7) {
    out.chain.push(level + ': ' + parent.tagName.toLowerCase() + (parent.id ? '#' + parent.id : '') + ' .' + cls(parent).slice(0, 120));
    let sib = child.nextElementSibling;
    while (sib) {
      out.following.push({level: level, tag: sib.tagName.toLowerCase(), cls: cls(sib).slice(0, 120),
                          html: sib.outerHTML.slice(0, 4000)});
      sib = sib.nextElementSibling;
    }
    child = parent; parent = parent.parentElement; level++;
  }
  return out;
}"""

ENUM_JS = "([css]) => {" + _FIND_TABLE + r"""
  const table = findTable();
  const vis = el => (el.offsetParent !== null || el.getClientRects().length > 0) && el.getBoundingClientRect().width > 0;
  const tr = table ? table.getBoundingClientRect() : null;
  const tableBottom = tr ? tr.bottom + window.scrollY : 0;
  const base = css + ", button, [role='button'], a[role='button'], input, textarea";
  const out = [];
  document.querySelectorAll(base).forEach(el => {
    if (!vis(el)) return;
    const r = el.getBoundingClientRect();
    const absTop = r.top + window.scrollY;
    const inTfoot = !!el.closest('tfoot');
    const inBody = table && table.contains(el) && !inTfoot;
    const afterTable = table ? ((table.compareDocumentPosition(el) & Node.DOCUMENT_POSITION_FOLLOWING) && !table.contains(el)) : true;
    if (inBody) return;
    if (!(inTfoot || afterTable || absTop >= tableBottom - 150)) return;
    out.push({
      tag: el.tagName.toLowerCase(), role: el.getAttribute('role') || '', type: el.getAttribute('type') || '',
      cls: ((el.className && el.className.toString) ? el.className.toString() : '').slice(0, 70),
      text: (el.innerText || el.value || '').replace(/\s+/g, ' ').trim().slice(0, 40),
      aria: el.getAttribute('aria-label') || '', title: el.getAttribute('title') || '',
      selectLike: el.matches(css), x: Math.round(r.left), y: Math.round(r.top),
      w: Math.round(r.width), h: Math.round(r.height), absTop: Math.round(absTop),
      cx: r.left + r.width / 2, cy: r.top + r.height / 2, left: r.left, top: r.top
    });
  });
  return out.slice(0, 90);
}"""

PAGINATION_JS = "([css]) => {" + _FIND_TABLE + r"""
  const table = findTable();
  const vis = el => (el.offsetParent !== null || el.getClientRects().length > 0) && el.getBoundingClientRect().width > 0;
  const cls = e => (e.className && e.className.toString) ? e.className.toString().trim() : '';
  const tableTop = table ? table.getBoundingClientRect().top + window.scrollY : 0;
  const below = el => (el.getBoundingClientRect().top + window.scrollY) >= tableTop - 5 && !(table && table.contains(el) && !el.closest('tfoot'));
  const named = Array.from(document.querySelectorAll(
    "nav, [role='navigation'], [class*='pagination' i], [class*='paginator' i], .v-pagination, .MuiPagination-root, " +
    ".MuiTablePagination-root, mat-paginator, .mat-mdc-paginator, .ant-pagination, .v-data-footer, " +
    ".v-data-table-footer, .q-table__bottom, tfoot")).filter(el => vis(el) && below(el));
  const isDigit = el => /^\d{1,5}$/.test((el.innerText || '').trim());
  const digits = Array.from(document.querySelectorAll('button, a, li, [role="button"]'))
      .filter(el => vis(el) && below(el) && isDigit(el));
  const common = list => {
    let a = list[0].parentElement;
    while (a && !list.every(x => a.contains(x))) a = a.parentElement;
    return a;
  };
  let root = null, how = '';
  if (named.length) { root = named[named.length - 1]; how = 'named container'; }
  else if (digits.length >= 2) { root = common(digits); how = 'common parent of ' + digits.length + ' page-number buttons'; }
  else if (digits.length === 1) { root = digits[0].parentElement; how = 'parent of the only page-number button'; }
  if (!root) return {root: null, digits: digits.length, named: named.length};
  root.scrollIntoView({block: 'center', behavior: 'instant'});
  const chain = [];
  let n = root;
  for (let i = 0; i < 6 && n && n !== document.body; i++) {
    chain.push(i + ': ' + n.tagName.toLowerCase() + (n.id ? '#' + n.id : '') + ' .' + cls(n).slice(0, 100));
    n = n.parentElement;
  }
  // look for a select-like control in the pagination root and then in its ancestors
  let control = null, level = -1;
  n = root;
  for (let i = 0; i < 6 && n && !control; i++) {
    const cands = Array.from(n.querySelectorAll(css)).filter(c => {
      if (!vis(c)) return false;
      const r = c.getBoundingClientRect();
      if (r.width < 30 || r.height < 16) return false;
      if (isDigit(c) && c.tagName !== 'SELECT') return false;
      return true;
    });
    cands.sort((a, b) => {
      const ra = a.getBoundingClientRect(), rb = b.getBoundingClientRect();
      return ra.width * ra.height - rb.width * rb.height;
    });
    if (cands.length) { control = cands[0]; level = i; }
    n = n.parentElement;
  }
  const rr = root.getBoundingClientRect();
  const res = {root: {tag: root.tagName.toLowerCase(), cls: cls(root).slice(0, 150), how: how,
                      html: root.outerHTML.slice(0, 2500), left: rr.left, top: rr.top, w: rr.width, h: rr.height},
               chain: chain, digits: digits.length, named: named.length, control: null, level: level};
  if (control) {
    const r = control.getBoundingClientRect();
    res.control = {tag: control.tagName.toLowerCase(), cls: cls(control).slice(0, 120),
                   left: r.left, top: r.top, w: r.width, h: r.height, text: (control.innerText || '').slice(0, 30)};
  }
  return res;
}"""

LABEL_JS = "([css]) => {" + _FIND_TABLE + r"""
  const table = findTable();
  const tableTop = table ? table.getBoundingClientRect().top + window.scrollY : 0;
  const vis = el => (el.offsetParent !== null || el.getClientRects().length > 0) && el.getBoundingClientRect().width > 0;
  const own = el => Array.from(el.childNodes).filter(n => n.nodeType === 3)
                         .map(n => n.textContent).join(' ').replace(/[*:]/g, '').replace(/\s+/g, ' ').trim().toLowerCase();
  const re = /^(items|items per page|item per page|rows per page|per page|show items|show rows|tampilkan|baris per halaman|item per halaman|jumlah baris|jumlah data)$/;
  let labels = Array.from(document.querySelectorAll('body *')).filter(el => vis(el) && re.test(own(el)));
  labels = labels.filter(el => (el.getBoundingClientRect().top + window.scrollY) >= tableTop - 5);
  if (!labels.length) return null;
  const label = labels[labels.length - 1];
  label.scrollIntoView({block: 'center', behavior: 'instant'});
  const lr = label.getBoundingClientRect();
  const lcy = (lr.top + lr.bottom) / 2;
  const cls = e => (e.className && e.className.toString) ? e.className.toString().trim() : '';
  let control = null, node = label, level = -1;
  for (let i = 0; i < 5 && !control; i++) {
    node = node.parentElement;
    if (!node) break;
    const cands = Array.from(node.querySelectorAll(css)).filter(c => {
      if (!vis(c) || c.contains(label) || label.contains(c)) return false;
      const r = c.getBoundingClientRect();
      const cy = (r.top + r.bottom) / 2;
      return r.width >= 30 && r.height >= 16 && r.right > lr.left && Math.abs(cy - lcy) < 45;
    });
    cands.sort((a, b) => {
      const ra = a.getBoundingClientRect(), rb = b.getBoundingClientRect();
      return ra.width * ra.height - rb.width * rb.height;
    });
    if (cands.length) { control = cands[0]; level = i; }
  }
  const res = {label: {text: own(label), tag: label.tagName.toLowerCase(), cls: cls(label).slice(0, 100),
                       left: lr.left, right: lr.right, top: lr.top, bottom: lr.bottom},
               control: null, level: level,
               html: (label.parentElement && label.parentElement.parentElement
                      ? label.parentElement.parentElement : label.parentElement || label).outerHTML.slice(0, 1500)};
  if (control) {
    const r = control.getBoundingClientRect();
    res.control = {tag: control.tagName.toLowerCase(), cls: cls(control).slice(0, 120),
                   left: r.left, top: r.top, w: r.width, h: r.height};
  }
  return res;
}"""

NATIVE_SELECT_JS = r"""() => {
  document.querySelectorAll('[data-pw-items]').forEach(e => e.removeAttribute('data-pw-items'));
  const sizes = ['10', '50', '100', '250', '500'];
  const all = Array.from(document.querySelectorAll('select'));
  const info = {total: all.length, matches: []};
  all.forEach((s, i) => {
    const vals = Array.from(s.options).map(o => (o.value || o.text || '').trim());
    const hit = sizes.filter(z => vals.includes(z) || Array.from(s.options).some(o => (o.text || '').trim() === z));
    if (hit.length >= 2) info.matches.push({i: i, vals: vals.slice(0, 12), value: s.value});
  });
  if (info.matches.length) {
    all[info.matches[info.matches.length - 1].i].setAttribute('data-pw-items', '1');
    all[info.matches[info.matches.length - 1].i].scrollIntoView({block: 'center', behavior: 'instant'});
  }
  return info;
}"""

TEXT_HITS_JS = "() => {" + _FIND_TABLE + r"""
  const table = findTable();
  const re = /items|per page|rows|tampilkan|halaman|show/i;
  const own = el => Array.from(el.childNodes).filter(n => n.nodeType === 3).map(n => n.textContent).join(' ').replace(/\s+/g, ' ').trim();
  const out = [];
  document.querySelectorAll('body *').forEach(el => {
    const t = own(el);
    if (!t || t.length > 40 || !re.test(t)) return;
    const r = el.getBoundingClientRect();
    out.push({tag: el.tagName.toLowerCase(), text: t, inTable: !!(table && table.contains(el)),
              inTfoot: !!el.closest('tfoot'),
              visible: (el.offsetParent !== null || el.getClientRects().length > 0) && r.width > 0,
              absTop: Math.round(r.top + window.scrollY), absLeft: Math.round(r.left)});
  });
  return {hits: out.slice(0, 30), docH: document.documentElement.scrollHeight, bodyTail: (document.body.innerText || '').slice(-500)};
}"""

# Footer range text such as "501-1000 of 128,345" -- anchored to the Closed Alarm table (the Daily Analytics
# tables have their own footers and must be ignored).
PAGE_RANGE_JS = "() => {" + _FIND_TABLE + r"""
  const table = findTable();
  if (!table) return null;
  const re = /(\d[\d.,]*)\s*[-\u2013\u2014]\s*(\d[\d.,]*)\s*(?:of|dari|\/)\s*(\d[\d.,]*)/i;
  const after = el => table.contains(el) ? !!el.closest('tfoot')
                                          : !!(table.compareDocumentPosition(el) & Node.DOCUMENT_POSITION_FOLLOWING);
  const sel = "tfoot, [class*='pagination' i], [class*='footer' i], mat-paginator, .MuiTablePagination-root, " +
              ".q-table__bottom, .ant-pagination";
  for (const el of document.querySelectorAll(sel)) {
    if (!after(el)) continue;
    const t = (el.innerText || '');
    if (t.length > 600) continue;
    const m = t.match(re);
    if (m) return [m[1], m[2], m[3]];
  }
  for (const el of document.querySelectorAll('body *')) {          // fallback: any small element below the table
    if (el.children.length > 3 || !after(el)) continue;
    const t = (el.innerText || '').trim();
    if (!t || t.length > 60) continue;
    const m = t.match(re);
    if (m) return [m[1], m[2], m[3]];
  }
  return null;
}"""

FIND_OPTIONS_JS = r"""([ax, ay]) => {
  const sizes = ['10', '50', '100', '250', '500'];
  const found = [];
  document.querySelectorAll('body *').forEach(el => {
    if (el.closest('table')) return;
    const r = el.getBoundingClientRect();
    if (r.width <= 0 || r.height <= 0) return;
    const s = getComputedStyle(el);
    if (s.visibility === 'hidden' || s.display === 'none') return;
    const t = (el.innerText || '').replace(/\s+/g, ' ').trim();
    if (!sizes.includes(t)) return;
    if (Math.abs(r.left - ax) > 500 || Math.abs((r.top + r.bottom) / 2 - ay) > 800) return;
    found.push({el, t, x: r.left + r.width / 2, y: r.top + r.height / 2, w: r.width, h: r.height});
  });
  const keep = found.filter(o => !found.some(p => p !== o && o.el.contains(p.el)));   // deepest only
  return keep.map(o => ({t: o.t, x: o.x, y: o.y, w: o.w, h: o.h}));
}"""


def _num(x) -> int:
    return int(re.sub(r"\D", "", str(x)) or 0)


def detect_site(page, job_dir: Path, verbose=True) -> dict:
    """Detect Vuetify / MUI / Angular Material / Ant / Quasar / AG Grid and whether the table is virtualized."""
    try:
        counts = page.evaluate(FRAMEWORK_JS, FRAMEWORK_SELECTORS)
    except PWError as exc:
        job_log(job_dir, f"SITE detection failed: {str(exc)[:150]}", "WARNING")
        return {}
    z = lambda k: isinstance(counts.get(k), int) and counts[k] > 0
    if z("vuetify_app") or z("vuetify_table") or z("vuetify2_footer") or z("vuetify3_footer") or z("vuetify_select"):
        fw = "vuetify"
    elif z("mui_table") or z("mui_pagination") or z("mui_datagrid"):
        fw = "mui"
    elif z("mat_paginator"):
        fw = "angular-material"
    elif z("ant_table") or z("ant_pagination"):
        fw = "antd"
    elif z("quasar"):
        fw = "quasar"
    elif z("ag_grid"):
        fw = "ag-grid"
    else:
        fw = "generic"
    virt = z("vuetify_virtual") or z("mui_virtual") or z("ag_grid") or z("virtual_strict")
    SITE_INFO["framework"], SITE_INFO["virtualized"] = fw, virt
    if verbose:
        nonzero = {k: v for k, v in counts.items() if v}
        job_log(job_dir, f"SITE: framework={fw} virtualized_table={virt} "
                         f"(loose 'virtual' class matches={counts.get('virtual_loose')}) counts={nonzero}")
        if virt:
            job_log(job_dir, "SITE: virtualized table detected -> DOM row counts can under-report; "
                             "the footer range text ('1-500 of N') is used as an extra signal.", "WARNING")
    return counts


def read_range(page):
    """(first, last, total) from the footer text below the Closed Alarm table, or None."""
    try:
        r = page.evaluate(PAGE_RANGE_JS)
    except PWError:
        return None
    if not r:
        return None
    a, b, t = (_num(x) for x in r)
    return (a, b, t) if b >= a > 0 else None


def effective_rows(page) -> int:
    """Rows on the page. For virtualized tables the DOM only holds the visible window, so use the footer range."""
    n = count_rows(page)
    if SITE_INFO.get("virtualized"):
        r = read_range(page)
        if r:
            n = max(n, r[1] - r[0] + 1)
    return n


def scroll_to_footer(page, job_dir: Path, label="", verbose=True, passes=3):
    """Scroll every scrollable ancestor of the table AND the window to the very bottom."""
    last = None
    for _ in range(passes):
        try:
            last = page.evaluate(SCROLL_FOOTER_JS)
        except PWError as exc:
            job_log(job_dir, f"[scroll {label}] JS scroll failed: {str(exc)[:150]}", "WARNING")
        page.wait_for_timeout(500)
    try:
        page.mouse.move(800, 600)
        page.mouse.wheel(0, 4000)
    except PWError:
        pass
    page.wait_for_timeout(700)
    if verbose and last:
        conts = "; ".join(f"{c['el']} scrollTop {c['before']}->{c['after']} (scrollH={c['scrollH']}, clientH={c['clientH']})"
                          for c in last.get("containers", [])) or "none"
        w = last.get("win") or {}
        job_log(job_dir, f"SCROLL [{label}] table_found={not last.get('noTable')} "
                         f"scrollable containers: {conts} | window scrollTop {w.get('before')}->{w.get('after')} "
                         f"(scrollH={w.get('scrollH')}, innerH={w.get('innerH')})")
    return last


def diagnose_footer(page, job_dir: Path, tag: str, title: str):
    """Dump footer HTML, enumerate controls near the bottom, find pagination root, list 'items'-like texts."""
    job_log(job_dir, f"--- DOM INSPECTION ({title}) ---")
    append_dom(job_dir, tag, f"\n===== {title} @ {datetime.now():%H:%M:%S} =====")
    detect_site(page, job_dir)

    try:
        dump = page.evaluate(FOOTER_DUMP_JS)
        if dump.get("error"):
            job_log(job_dir, f"FOOTER DUMP: {dump['error']}", "WARNING")
            append_dom(job_dir, tag, f"FOOTER DUMP: {dump['error']}")
        else:
            job_log(job_dir, "FOOTER ancestors: " + " | ".join(dump["chain"]))
            job_log(job_dir, f"FOOTER <tfoot> html: {dump['tfoot'][:600] or '(none)'}")
            append_dom(job_dir, tag, "ANCESTORS:\n" + "\n".join(dump["chain"]))
            append_dom(job_dir, tag, "TFOOT:\n" + (dump["tfoot"] or "(none)"))
            if not dump["following"]:
                job_log(job_dir, "FOOTER: the table has NO following siblings at any of the 7 ancestor levels.",
                        "WARNING")
            for item in dump["following"]:
                job_log(job_dir, f"FOOTER sibling (level {item['level']}) <{item['tag']} class='{item['cls']}'> "
                                 f"{item['html'][:500]}")
                append_dom(job_dir, tag, f"SIBLING level {item['level']} <{item['tag']} class='{item['cls']}'>\n{item['html']}")
    except PWError as exc:
        job_log(job_dir, f"FOOTER DUMP failed: {str(exc)[:150]}", "WARNING")

    try:
        els = page.evaluate(ENUM_JS, [SELECT_LIKE_CSS])
        job_log(job_dir, f"ENUM: {len(els)} visible select/combobox/button/dropdown/input elements near the bottom")
        for e in els:
            line = (f"  <{e['tag']}{'[' + e['type'] + ']' if e['type'] else ''}> role={e['role'] or '-'} "
                    f"selectLike={e['selectLike']} text='{e['text']}' aria='{e['aria']}' class='{e['cls']}' "
                    f"at x={e['x']} y={e['y']} {e['w']}x{e['h']} absTop={e['absTop']}")
            job_log(job_dir, line)
            append_dom(job_dir, tag, "ENUM" + line)
    except PWError as exc:
        job_log(job_dir, f"ENUM failed: {str(exc)[:150]}", "WARNING")

    try:
        pag = page.evaluate(PAGINATION_JS, [SELECT_LIKE_CSS])
        if not pag.get("root"):
            job_log(job_dir, f"PAGINATION: no pagination container found (page-number buttons={pag.get('digits')}, "
                             f"named containers={pag.get('named')}).", "WARNING")
        else:
            r = pag["root"]
            job_log(job_dir, f"PAGINATION: root=<{r['tag']} class='{r['cls']}'> via {r['how']}; "
                             f"{pag['digits']} page-number buttons; chain: {' | '.join(pag['chain'])}")
            job_log(job_dir, f"PAGINATION root html: {r['html'][:700]}")
            append_dom(job_dir, tag, "PAGINATION ROOT:\n" + r["html"] + "\nCHAIN:\n" + "\n".join(pag["chain"]))
            if pag.get("control"):
                job_log(job_dir, f"PAGINATION: select-like control found {pag['control']} at ancestor level {pag['level']}")
    except PWError as exc:
        job_log(job_dir, f"PAGINATION probe failed: {str(exc)[:150]}", "WARNING")

    try:
        th = page.evaluate(TEXT_HITS_JS)
        job_log(job_dir, f"TEXT HITS ('items'/'per page'/'rows'/...): {len(th['hits'])}, document height={th['docH']}")
        for h in th["hits"]:
            job_log(job_dir, f"  '{h['text']}' <{h['tag']}> inTable={h['inTable']} inTfoot={h['inTfoot']} "
                             f"visible={h['visible']} absTop={h['absTop']} absLeft={h['absLeft']}")
        job_log(job_dir, f"BODY TAIL (last 500 chars of page text): {th['bodyTail']!r}")
        append_dom(job_dir, tag, "TEXT HITS: " + json.dumps(th["hits"]) + "\nBODY TAIL: " + th["bodyTail"])
    except PWError as exc:
        job_log(job_dir, f"TEXT HITS failed: {str(exc)[:150]}", "WARNING")


def _point(how, left, top, w, h):
    return {"kind": "point", "how": how, "left": left, "right": left + w, "top": top, "bottom": top + h,
            "w": w, "h": h, "cx": left + w / 2, "cy": top + h / 2}


def locate_control_once(page, job_dir: Path):
    """
    Try every strategy once and log each attempt (INFO lines are silenced when SITE_INFO['quiet']). Returns
      {'kind': 'native', ...}  for a real <select> that contains the sizes, or
      {'kind': 'point', cx, cy, ...} for a custom dropdown (click by coordinates), or None.
    """
    try:
        info = page.evaluate(NATIVE_SELECT_JS)
        if info["matches"]:
            m = info["matches"][-1]
            llog(job_dir, f"LOCATE [native <select>] -> FOUND options={m['vals']} current='{m['value']}' "
                          f"({info['total']} selects on page)")
            return {"kind": "native", "how": "native select"}
        llog(job_dir, f"LOCATE [native <select>] -> none ({info['total']} <select> elements, none with page sizes)")
    except PWError as exc:
        llog(job_dir, f"LOCATE [native <select>] -> error {str(exc)[:120]}", "WARNING")

    for name, css in ITEMS_SELECTORS:
        try:
            loc = page.locator(css)
            n = loc.count()
            vis = [i for i in range(min(n, 8)) if loc.nth(i).is_visible()]
            llog(job_dir, f"LOCATE [{name}] css='{css[:80]}' -> {n} match(es), {len(vis)} visible")
            if vis:
                el = loc.nth(vis[-1])
                el.scroll_into_view_if_needed(timeout=3000)
                box = el.bounding_box()
                if box:
                    llog(job_dir, f"LOCATE [{name}] -> USING it at x={box['x']:.0f} y={box['y']:.0f} "
                                  f"{box['width']:.0f}x{box['height']:.0f}")
                    return _point(name, box["x"], box["y"], box["width"], box["height"])
        except PWError as exc:
            llog(job_dir, f"LOCATE [{name}] -> error {str(exc)[:120]}", "WARNING")

    try:
        res = page.evaluate(LABEL_JS, [SELECT_LIKE_CSS])
        if not res:
            llog(job_dir, "LOCATE [text label] -> no 'Items'/'per page'/'Rows per page'/'Tampilkan' label below the table")
        else:
            lab = res["label"]
            llog(job_dir, f"LOCATE [text label] -> label '{lab['text']}' <{lab['tag']} class='{lab['cls']}'> "
                          f"at x={lab['left']:.0f} y={lab['top']:.0f}; html: {res['html'][:500]}")
            if res.get("control"):
                c = res["control"]
                llog(job_dir, f"LOCATE [text label] -> USING select-like <{c['tag']} class='{c['cls']}'> "
                              f"(found {res['level']} level(s) above the label)")
                return _point("label+control", c["left"], c["top"], c["w"], c["h"])
            llog(job_dir, "LOCATE [text label] -> no select-like element next to the label; "
                          "will click to the right of the label", "WARNING")
            h = max(20.0, lab["bottom"] - lab["top"])
            return _point("label-offset", lab["right"] + 10, lab["top"], 70, h)
    except PWError as exc:
        llog(job_dir, f"LOCATE [text label] -> error {str(exc)[:120]}", "WARNING")

    try:
        pag = page.evaluate(PAGINATION_JS, [SELECT_LIKE_CSS])
        if not pag.get("root"):
            llog(job_dir, f"LOCATE [pagination-first] -> no pagination container "
                          f"(page-number buttons={pag.get('digits')}, named containers={pag.get('named')})")
        else:
            llog(job_dir, f"LOCATE [pagination-first] -> pagination root <{pag['root']['tag']} "
                          f"class='{pag['root']['cls']}'> via {pag['root']['how']}")
            if pag.get("control"):
                c = pag["control"]
                llog(job_dir, f"LOCATE [pagination-first] -> USING select-like <{c['tag']} class='{c['cls']}'> "
                              f"text='{c['text']}' (ancestor level {pag['level']})")
                return _point("pagination+control", c["left"], c["top"], c["w"], c["h"])
            llog(job_dir, "LOCATE [pagination-first] -> no select-like control inside the pagination container "
                          "or its 5 ancestors", "WARNING")
    except PWError as exc:
        llog(job_dir, f"LOCATE [pagination-first] -> error {str(exc)[:120]}", "WARNING")

    try:
        els = page.evaluate(ENUM_JS, [SELECT_LIKE_CSS])
        hits = [e for e in els if e["selectLike"] and e["w"] >= 30 and e["h"] >= 16
                and re.fullmatch(r"(10|50|100|250|500)", re.sub(r"\D", "", e["text"]) or "x") and len(e["text"]) <= 12]
        llog(job_dir, f"LOCATE [shows-a-page-size] -> {len(els)} candidates near the bottom, {len(hits)} look like "
                      f"a size selector")
        if hits:
            e = hits[-1]
            llog(job_dir, f"LOCATE [shows-a-page-size] -> USING <{e['tag']} class='{e['cls']}'> text='{e['text']}'")
            return _point("shows-a-page-size", e["left"], e["top"], e["w"], e["h"])
    except PWError as exc:
        llog(job_dir, f"LOCATE [shows-a-page-size] -> error {str(exc)[:120]}", "WARNING")
    return None


def find_items_control(page, job_dir: Path, tag="", rounds=4):
    """Scroll to the footer, screenshot, inspect the DOM, locate the Items control (retrying a few rounds)."""
    take_shot(page, job_dir, f"shot_{tag}before_scroll")
    for rnd in range(1, rounds + 1):
        scroll_to_footer(page, job_dir, f"{tag}round{rnd}")
        if rnd == 1:
            take_shot(page, job_dir, f"shot_{tag}after_scroll")
            diagnose_footer(page, job_dir, tag, "after first footer scroll")
        ctl = locate_control_once(page, job_dir)
        if ctl:
            job_log(job_dir, f"Items control found in round {rnd} via [{ctl['how']}].")
            return ctl
        job_log(job_dir, f"Round {rnd}/{rounds}: Items control not found yet; waiting 5s (the footer may render "
                         f"later than the table), then scrolling again.", "WARNING")
        page.wait_for_timeout(5000)
    scroll_to_footer(page, job_dir, f"{tag}final")
    take_shot(page, job_dir, f"shot_{tag}after_scroll_final")
    diagnose_footer(page, job_dir, tag, "FINAL inspection before giving up")
    return locate_control_once(page, job_dir)


def wait_footer_ready(page, job_dir: Path, timeout=60) -> bool:
    """After opening CLOSED ALARM: wait (scrolling to the bottom each time) until ANY footer hint shows up."""
    t0 = time.time()
    end = t0 + timeout
    last_log = 0.0
    while time.time() < end:
        scroll_to_footer(page, job_dir, "wait-footer", verbose=False, passes=1)
        try:
            pag = page.evaluate(PAGINATION_JS, [SELECT_LIKE_CSS])
            lab = page.evaluate(LABEL_JS, [SELECT_LIKE_CSS])
            nat = page.evaluate(NATIVE_SELECT_JS)
        except PWError:
            pag, lab, nat = {}, None, {"matches": []}
        if (pag and pag.get("root")) or lab or nat.get("matches"):
            job_log(job_dir, f"Footer is rendered after {int(time.time() - t0)}s "
                             f"(pagination root={'yes' if pag and pag.get('root') else 'no'}, "
                             f"items label={'yes' if lab else 'no'}, native select={'yes' if nat.get('matches') else 'no'}).")
            return True
        if time.time() - last_log >= 10:
            last_log = time.time()
            job_log(job_dir, f"Waiting for the pagination footer to render... {int(time.time() - t0)}s")
        page.wait_for_timeout(1000)
    job_log(job_dir, f"No pagination footer detected after {timeout}s; the full inspection will run next.", "WARNING")
    return False


def choose_size(page, job_dir: Path, size: str, tag: str = ""):
    """Locate the Items control again (page may have re-rendered), open it and choose `size`."""
    try:
        page.keyboard.press("Escape")
    except PWError:
        pass
    scroll_to_footer(page, job_dir, f"{tag}pick{size}", verbose=False, passes=2)
    ctl = locate_control_once(page, job_dir)
    if not ctl:
        raise RuntimeError("the Items control disappeared before selecting the size")

    if ctl["kind"] == "native":
        loc = page.locator("select[data-pw-items]").first
        for kw in ({"label": size}, {"value": size}):
            try:
                loc.select_option(force=True, timeout=8000, **kw)
                llog(job_dir, f"Native <select>: chose {kw}")
                return True
            except PWError as exc:
                llog(job_dir, f"Native select_option {kw} failed: {str(exc)[:150]}", "WARNING")
        raise RuntimeError("native select_option failed")

    offsets = [0.0]
    if ctl["w"] > 60:
        offsets += [ctl["w"] / 2 - 16, -ctl["w"] / 4]
    opts, texts = [], []
    for dx in offsets:
        page.mouse.click(ctl["cx"] + dx, ctl["cy"])
        page.wait_for_timeout(900)
        opts = page.evaluate(FIND_OPTIONS_JS, [ctl["left"], ctl["cy"]])
        texts = sorted({o["t"] for o in opts}, key=lambda s: int(s))
        llog(job_dir, f"Clicked the Items control (via [{ctl['how']}]) at dx={dx:+.0f}px, "
                      f"visible size entries: {texts}")
        if len(texts) >= 3 or (size in texts and len(texts) >= 2):
            break
        try:
            page.keyboard.press("Escape")
        except PWError:
            pass
        page.wait_for_timeout(300)
    if not SITE_INFO.get("quiet"):
        take_shot(page, job_dir, f"shot_{tag}dropdown_open_{size}")

    wanted = [o for o in opts if o["t"] == size]
    if wanted:
        o = sorted(wanted, key=lambda o: o["w"] * o["h"])[0]
        llog(job_dir, f"Clicking '{size}' at x={o['x']:.0f} y={o['y']:.0f}")
        page.mouse.click(o["x"], o["y"])
        return True

    opt_xpath = ("xpath=//*[(self::li or self::div or self::span or self::a or self::mat-option or self::option "
                 f"or @role='option' or @role='menuitem') and normalize-space(.)='{size}' and not(ancestor::table)]")
    try:
        cand = first_visible(page, opt_xpath, 4)
        llog(job_dir, f"Option '{size}' found by text; clicking it.")
        safe_click(cand)
        return True
    except (PWError, RuntimeError) as exc:
        llog(job_dir, f"Option-by-text lookup for '{size}' failed: {str(exc)[:120]}", "WARNING")

    idx = SIZE_ORDER.index(size)
    llog(job_dir, f"Entry '{size}' not visible; using the keyboard (ArrowDown x{idx} + Enter).", "WARNING")
    page.mouse.click(ctl["cx"], ctl["cy"])
    page.wait_for_timeout(500)
    for _ in range(idx):
        page.keyboard.press("ArrowDown")
        page.wait_for_timeout(120)
    page.keyboard.press("Enter")
    return True


def wait_rows_after_size(page, job_dir: Path, rows_before: int, size: int,
                         timeout=PAGE_SIZE_LOAD_TIMEOUT, rep=None, base_progress=0.0) -> int:
    """PATIENT wait after choosing a bigger page size (full size, or a stable partial table)."""
    t0 = time.time()
    end = t0 + timeout
    last_n, stable_since, last_log = -1, None, 0.0
    while time.time() < end:
        n = effective_rows(page)
        now = time.time()
        if n > rows_before:
            if n >= size:
                page.wait_for_timeout(1500)       # let the last rows render
                return max(n, effective_rows(page))
            if n == last_n:
                if stable_since is None:
                    stable_since = now
                elif now - stable_since >= PAGE_SIZE_STABLE_SECONDS:
                    return n
            else:
                stable_since = None
        else:
            stable_since = None
        last_n = n
        if now - last_log >= 10:
            last_log = now
            msg = f"Waiting for the table to show {size} rows... now {n} rows ({int(now - t0)}s)"
            job_log(job_dir, msg)
            if rep is not None:
                rep.running(msg, base_progress, quiet=True)
        page.wait_for_timeout(500)
    n = effective_rows(page)
    return n if n > rows_before else rows_before


def pw_set_page_size(page, job_dir: Path, rep=None, strict=True, tag="", base_progress=0.0, sizes=None) -> int:
    """FULL version (with footer inspection / screenshots): find the Items control, choose the size and WAIT for the rows."""
    sizes = tuple(sizes) if sizes else (str(TARGET_PAGE_SIZE),)
    detect_site(page, job_dir)
    n0 = effective_rows(page)
    if n0 and n0 < DEFAULT_PAGE_ROWS and not SITE_INFO["virtualized"]:
        job_log(job_dir, f"Only {n0} rows in the table: everything fits on one page, page size left as is.")
        return n0

    if rep is not None:
        rep.running("Scrolling to the table footer and looking for the Items dropdown...", base_progress, quiet=True)
    ctl = find_items_control(page, job_dir, tag)
    if ctl is None:
        take_shot(page, job_dir, f"shot_{tag}items_not_found")
        msg = ("Could not find the Items-per-page control even after scrolling to the bottom of the table and "
               "inspecting the DOM. Open Diagnostics: see job.log (lines starting with FOOTER / ENUM / PAGINATION / "
               "LOCATE / TEXT HITS), the dom_*diag.txt file and the shot_*scroll* screenshots.")
        if strict:
            raise RuntimeError(msg)
        job_log(job_dir, msg, "WARNING")
        return n0

    for size in sizes:
        try:
            choose_size(page, job_dir, size, tag)
            job_log(job_dir, f"Selected '{size}' in the Items dropdown, waiting for the table to reload "
                             f"(this can take a while)...")
            if rep is not None:
                rep.running(f"Selected {size} items per page, waiting for the table to load...",
                            base_progress, quiet=True)
            rows_now = wait_rows_after_size(page, job_dir, n0, int(size), rep=rep, base_progress=base_progress)
            wait_loading_done(page)
            page.wait_for_timeout(500)
            rows_now = max(rows_now, effective_rows(page))
            if rows_now > n0:
                job_log(job_dir, f"Page size set to {size}: {rows_now} rows on the page now.")
                return rows_now
            job_log(job_dir, f"Page size {size} did not change the table ({rows_now} rows).", "WARNING")
        except (PWError, RuntimeError) as exc:
            job_log(job_dir, f"Page size {size} failed: {str(exc)[:200]}", "WARNING")
        try:
            page.keyboard.press("Escape")
        except PWError:
            pass
    take_shot(page, job_dir, f"shot_{tag}size_not_applied")
    if strict:
        raise RuntimeError("The Items control was found but the table never grew above 10 rows (see the "
                           "shot_*dropdown_open_* screenshots and the LOCATE / Clicked lines in the log). "
                           "Stopping instead of downloading 10 rows at a time.")
    job_log(job_dir, "Page size could not be raised again (the result may genuinely fit in 10 rows).", "WARNING")
    return n0


# =====================================================================
#  PAGE-SIZE VERIFICATION (the size is never assumed: it is checked before Apply, after Apply,
#  after navigation and right before every download)
# =====================================================================
def page_size_ok(page, size: int) -> bool:
    """
    True only when the table really shows the wanted size:
      * footer range available -> rows on screen >= min(size, rows remaining from the first row of this page)
      * no footer range        -> more than 10 rows are visible
    Never True while only 10 rows are visible and more records exist.
    """
    n = effective_rows(page)
    if n <= 0:
        return False
    r = read_range(page)
    if r:
        expected = min(size, r[2] - r[0] + 1)
        return n >= expected
    return n > DEFAULT_PAGE_ROWS


def current_page(page, size: int) -> int:
    """Page number currently displayed: from the footer range ('501-1000 of N') or the active page button."""
    r = read_range(page)
    if r and size:
        return (r[0] - 1) // size + 1
    try:
        res = page.evaluate(MARK_PAGE_JS, 0)
        return int(res.get("cur") or 0)
    except (PWError, TypeError, ValueError):
        return 0


def log_state(page, job_dir: Path, size: int, tag: str):
    """One log line with everything that matters: size, visible rows, footer range, current page."""
    rows = effective_rows(page)
    rng = read_range(page)
    job_log(job_dir, f"STATE [{tag}] page_size_wanted={size} visible_rows={rows} footer_range={rng} "
                     f"current_page={current_page(page, size)} size_ok={page_size_ok(page, size)}")


def quick_set_size(page, job_dir: Path, size: int, rep=None, base_progress=0.0) -> int:
    """
    LIGHT version used in every per-page cycle: no screenshots, no DOM dumps, INFO logging silenced.
    Chooses exactly `size`. If that fails the FULL version runs once as a fallback (it logs everything).
    Returns the rows now shown.
    """
    n0 = effective_rows(page)
    if page_size_ok(page, size):
        return n0
    SITE_INFO["quiet"] = True
    try:
        for attempt in (1, 2):
            try:
                choose_size(page, job_dir, str(size), "q_")
                rows_now = wait_rows_after_size(page, job_dir, n0, size, rep=rep, base_progress=base_progress)
                wait_loading_done(page)
                page.wait_for_timeout(300)
                rows_now = max(rows_now, effective_rows(page))
                if rows_now > n0:
                    return rows_now
                job_log(job_dir, f"Quick page-size attempt {attempt}: table still shows {rows_now} rows.", "WARNING")
            except (PWError, RuntimeError) as exc:
                job_log(job_dir, f"Quick page-size attempt {attempt} failed: {str(exc)[:160]}", "WARNING")
            try:
                page.keyboard.press("Escape")
            except PWError:
                pass
    finally:
        SITE_INFO["quiet"] = False
    job_log(job_dir, "Quick page-size selection failed; running the full version with diagnostics.", "WARNING")
    return pw_set_page_size(page, job_dir, rep=rep, strict=True, tag="fb_", base_progress=base_progress,
                            sizes=(str(size),))


def ensure_page_size(page, job_dir: Path, state: dict, rep, stage: str, prog: float = 0.0) -> int:
    """
    Make sure the table shows TARGET_PAGE_SIZE rows (or all remaining rows when fewer exist).
    Waits for a slow reload first; if the size was reset (10 rows) it is selected again, up to
    SIZE_RESTORE_ATTEMPTS times. Raises when it cannot be restored: a cycle must never continue with 10 rows.
    """
    size = state["page_size"]
    for attempt in range(1, SIZE_RESTORE_ATTEMPTS + 1):
        end = time.time() + SIZE_VERIFY_WAIT
        while time.time() < end and not page_size_ok(page, size):
            page.wait_for_timeout(500)
        if page_size_ok(page, size):
            log_state(page, job_dir, size, f"{stage} -> OK")
            return effective_rows(page)

        log_state(page, job_dir, size, f"{stage} -> NOT {size}, restoring (attempt {attempt})")
        if not state.get("size_discovered"):
            # very first time in this job: full discovery of the Items control (screenshots + DOM dump)
            pw_set_page_size(page, job_dir, rep=rep, strict=True, tag="2a_", base_progress=prog,
                             sizes=(str(size),))
            state["size_discovered"] = True
            if state.get("session_no") == 1:
                take_shot(page, job_dir, "shot_1_after_page_size")
        else:
            quick_set_size(page, job_dir, size, rep=rep, base_progress=prog)
        wait_loading_done(page, 15000)
        page.wait_for_timeout(500)

    log_state(page, job_dir, size, f"{stage} -> FAILED")
    raise RuntimeError(f"[{stage}] the page size could not be set to {size} "
                       f"(table shows {effective_rows(page)} rows). Not continuing with a small page.")


# =====================================================================
#  Apply filter, page navigation, page cycle
# =====================================================================
def apply_filters(page, job_dir: Path, expect_change: bool, size: int = TARGET_PAGE_SIZE):
    """Click Apply, wait for the reload, the loading indicators and a stable table. Logs the size before/after."""
    log_state(page, job_dir, size, "BEFORE Apply")
    old = get_sig(page)
    pw_click_apply(page)
    changed = wait_sig_change(page, old, APPLY_LOAD_TIMEOUT if expect_change else APPLY_QUICK_TIMEOUT)
    wait_loading_done(page, 30000)
    stable = wait_table_stable(page)
    page.wait_for_timeout(500)
    job_log(job_dir, f"Apply done: table_changed={changed} stable={stable} rows={effective_rows(page)}")
    log_state(page, job_dir, size, "AFTER Apply (before size check)")


def size_was_reset(page, size: int, target: int, state: dict) -> bool:
    """True when the table shows <= 10 rows although a bigger page size was chosen and more rows remain."""
    if size <= DEFAULT_PAGE_ROWS:
        return False
    n = effective_rows(page)
    if n == 0 or n > DEFAULT_PAGE_ROWS:
        return False
    r = read_range(page)
    if r:
        return (r[2] - r[0] + 1) > n          # more rows remain than are displayed
    return target < state["total_pages"]


def goto_page(page, job_dir: Path, target: int, state: dict, rep, monitor, browser):
    """Go to page `target`; the page size is re-verified on arrival and restored to 500 whenever it was reset."""
    size = state["page_size"]
    resets = 0
    for step in range(1, MAX_NAV_STEPS + 1):
        cur = current_page(page, size)

        if cur == target:
            if page_size_ok(page, size):
                log_state(page, job_dir, size, f"NAV page {target} reached after {step - 1} click(s)")
                return
            resets += 1
            job_log(job_dir, f"NAV: on page {target} but the size is NOT {size} (reset #{resets}); restoring.",
                    "WARNING")
            if resets > 4:
                raise RuntimeError(f"The page size keeps resetting while navigating to page {target}.")
            ensure_page_size(page, job_dir, state, rep, f"nav-restore p{target}")
            continue                          # a restore may jump back to page 1 -> the loop navigates again

        try:
            res = page.evaluate(MARK_PAGE_JS, target)
        except PWError as exc:
            raise RuntimeError(f"Could not read the page buttons: {str(exc)[:120]}")
        mode = res.get("mode")
        old = get_sig(page)
        if mode in ("exact", "jump"):
            safe_click(page.locator("[data-pw-page]").first)
        else:
            outcome = click_next(page)
            if outcome != "clicked":
                raise RuntimeError(f"Cannot reach page {target}: no page button visible and 'Next' is "
                                   f"'{outcome}' (current page {cur}, visible buttons {res.get('digits')}).")
        changed = wait_sig_change(page, old, 40)
        wait_loading_done(page, 15000)
        page.wait_for_timeout(500)
        job_log(job_dir, f"NAV step {step}: target={target} was_on={cur} click={mode if mode != 'none' else 'next'} "
                         f"table_changed={changed} rows_now={effective_rows(page)} buttons={res.get('digits')} "
                         f"footer_range={read_range(page)}")
        if monitor.tripped:
            raise RuntimeError(monitor.tripped)
        if not browser.is_connected() or page.is_closed():
            raise RuntimeError("Chromium exited while navigating.")

        if size_was_reset(page, size, target, state):
            resets += 1
            job_log(job_dir, f"NAV: the page size was RESET to {effective_rows(page)} rows after the click "
                             f"(reset #{resets}); selecting {size} again.", "WARNING")
            if resets > 4:
                raise RuntimeError(f"The page size keeps resetting to 10 while navigating to page {target}; "
                                   f"cannot reach it.")
            ensure_page_size(page, job_dir, state, rep, f"nav-reset p{target}")
            continue

        if mode == "exact" and changed and current_page(page, size) == 0:
            job_log(job_dir, f"NAV: clicked page button {target}; the current page cannot be read, assuming success.")
            return
    raise RuntimeError(f"Could not reach page {target} within {MAX_NAV_STEPS} clicks.")


def click_next(page) -> str:
    """Click the 'next page' button. Returns 'clicked' | 'disabled' | 'none'."""
    el, disabled = find_next(page)
    if el is None:
        return "none"
    try:
        if disabled:
            return "disabled"
        safe_click(el)
        return "clicked"
    finally:
        _dispose(el)


def find_next(page):
    """Returns (element_handle_or_None, disabled). The CALLER must dispose the returned handle."""
    handle = page.evaluate_handle(FIND_NEXT_JS)
    el = handle.as_element()
    if el is None:
        _dispose(handle)
        return None, True
    try:
        return el, bool(el.evaluate(IS_DISABLED_JS))
    except PWError:
        _dispose(el)
        return None, True


# =====================================================================
#  Download Alarm > By Page  (one file per page)
# =====================================================================
BY_PAGE_XPATH = (f"xpath=//*[(self::button or self::a or self::li or self::span or self::div "
                 f"or @role='menuitem') and {UPPER}='BY PAGE']")


def file_sha256(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def pw_download_by_page(page, job_dir: Path, dest_base: Path) -> Path:
    """Open 'Download Alarm', click 'By Page', save the file, then delete Playwright's own copy."""
    last_exc = None
    for attempt in (1, 2, 3):
        try:
            try:
                page.keyboard.press("Escape")  # make sure no old menu is open
            except PWError:
                pass
            with page.expect_download(timeout=DOWNLOAD_TIMEOUT * 1000) as info:
                safe_click(first_visible(page, f"xpath=//button[contains({UPPER}, 'DOWNLOAD ALARM')]", 20))
                page.wait_for_timeout(600)
                safe_click(first_visible(page, BY_PAGE_XPATH, 15))
            dl = info.value
            failure = dl.failure()
            if failure:
                raise RuntimeError(f"browser download failed: {failure}")
            ext = Path(dl.suggested_filename or "").suffix.lower() or ".xlsx"
            dest = Path(f"{dest_base}{ext}")
            dl.save_as(str(dest))
            try:
                dl.delete()          # drop Playwright's duplicate copy in downloads_path
            except Exception:
                pass
            if not dest.exists() or dest.stat().st_size == 0:
                raise RuntimeError("the downloaded file is empty")
            return dest
        except (PWError, RuntimeError) as exc:
            last_exc = exc
            job_log(job_dir, f"By Page download attempt {attempt} failed: {str(exc)[:200]}", "WARNING")
            page.wait_for_timeout(2000)
    raise RuntimeError(f"Download Alarm > By Page failed 3 times: {str(last_exc)[:200]}")


def probe_xlsx_rows(path: Path) -> int:
    """Real data-row count inside a downloaded XLSX (streaming, header excluded). -1 = unknown."""
    if path.suffix.lower() != ".xlsx":
        return -1
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:
        return -1
    try:
        n = 0
        for row in wb.worksheets[0].iter_rows(values_only=True):
            if any(c not in (None, "") for c in row):
                n += 1
        return max(0, n - 1)
    except Exception:
        return -1
    finally:
        wb.close()


def recycle_reason(snap: dict, baseline, session_pages: int):
    """Decide (between two pages) whether the browser should be closed and relaunched."""
    if RECYCLE_EVERY_PAGES and session_pages >= RECYCLE_EVERY_PAGES:
        return f"scheduled recycle after {session_pages} pages"
    if session_pages < RECYCLE_MIN_PAGES:
        return None
    limit = snap.get("limit")
    ws = snap.get("ws")
    if limit and ws and ws > limit * SOFT_RECYCLE_FRACTION:
        return f"working set {ws:.0f} MB is above {SOFT_RECYCLE_FRACTION:.0%} of the {limit:.0f} MB limit"
    if baseline and snap["chrom_rss"] > baseline + CHROMIUM_GROWTH_MB:
        return (f"Chromium RSS grew from {baseline:.0f} MB to {snap['chrom_rss']:.0f} MB "
                f"(+{CHROMIUM_GROWTH_MB} MB allowed)")
    return None


def fetch_page(page, job_dir: Path, dl_dir: Path, pages_dir: Path, rep: StepReporter, monitor, browser,
               state: dict, n: int, start_d: date, end_d: date, first_cycle: bool) -> bool:
    """
    ONE page cycle:
      1 size 500 -> 2 Start Date via the calendar -> 3 End Date untouched -> 4 Apply -> 5 verify 500 again
      -> 6 go to page n and verify 500 again -> 7 Download Alarm > By Page.
    Returns True when page n was the LAST page.
    """
    size = state["page_size"]
    total = max(state["total_pages"], n)
    prog = 0.46 + 0.44 * min((n - 1) / max(total, 1), 1.0)
    rep.running(f"Page {n} of ~{total}: 500 rows, start date, Apply, go to page, download...", prog, quiet=True)
    job_log(job_dir, f"P{n} CYCLE start | {format_snapshot(collect_snapshot(dl_dir))}")

    # ---- STEP 1: items per page FIRST
    ensure_page_size(page, job_dir, state, rep, f"P{n} step1 before dates", prog)

    # ---- STEP 2 + 3: Start Date through the calendar, End Date untouched
    ensure_dates(page, job_dir, start_d, end_d)

    # ---- STEP 4: Apply
    apply_filters(page, job_dir, expect_change=first_cycle, size=size)

    # ---- STEP 5: the site may reset the size after Apply
    ensure_page_size(page, job_dir, state, rep, f"P{n} step5 after Apply", prog)

    # ---- STEP 6: navigate (after Apply the table is on page 1), then verify the size once more
    if n > 1:
        goto_page(page, job_dir, n, state, rep, monitor, browser)
    ensure_page_size(page, job_dir, state, rep, f"P{n} step6 after navigation", prog)

    # ---- final gate before the download
    rows = effective_rows(page)
    rng = read_range(page)
    cur = current_page(page, size)
    job_log(job_dir, f"P{n} READY visible_rows={rows} footer_range={rng} current_page={cur} page_size={size}")
    if rng and rng[0] != (n - 1) * size + 1:
        raise RuntimeError(f"Wrong page on screen: the footer range starts at row {rng[0]} but page {n} "
                           f"must start at row {(n - 1) * size + 1}.")
    if cur not in (0, n):
        raise RuntimeError(f"Wrong page on screen: the current page is {cur}, expected {n}.")
    if not page_size_ok(page, size):
        raise RuntimeError(f"Refusing to download page {n}: the page size is not {size} ({rows} rows visible).")

    # ---- STEP 7: download
    job_log(job_dir, f"P{n} BEFORE download | {format_snapshot(collect_snapshot(dl_dir))}")
    dest = pw_download_by_page(page, job_dir, pages_dir / f"page_{n:04d}")      # raises if missing / empty
    job_log(job_dir, f"P{n} DOWNLOAD OK: {dest.name} {dest.stat().st_size / 1024:.0f} KB "
                     f"(exists={dest.exists()}, non-empty={dest.stat().st_size > 0})")
    job_log(job_dir, f"P{n} AFTER download  | {format_snapshot(collect_snapshot(dl_dir))}")

    digest = file_sha256(dest)
    if digest in state["hashes"]:
        dest.unlink(missing_ok=True)
        raise RuntimeError(f"The file downloaded for page {n} is identical to page {state['hashes'][digest]} "
                           f"(the table was not on page {n}).")
    state["hashes"][digest] = n
    state["files"].append(dest)
    state["resume_page"] = n + 1

    if n <= PROBE_FIRST_PAGES or n % PROBE_EVERY == 0:
        n_file = probe_xlsx_rows(dest)
        n_scr = count_rows(page)
        level = "WARNING" if n_file not in (-1, n_scr) else "INFO"
        job_log(job_dir, f"P{n} VERIFY file_rows={n_file} screen_rows={n_scr} size={dest.stat().st_size / 1024:.0f}KB",
                level)

    # ---- is this the last page?
    last = False
    rng = read_range(page)
    try:
        if rng and size:
            state["total_pages"] = max(state["total_pages"], n, math.ceil(rng[2] / size))
        else:
            state["total_pages"] = max(state["total_pages"], n, int(page.evaluate(TOTAL_PAGES_JS) or 0))
    except PWError:
        pass
    if rng and rng[1] >= rng[2]:
        last = True
        job_log(job_dir, f"Page {n}: footer range {rng[0]}-{rng[1]} of {rng[2]} -> this is the last page.")
    else:
        nxt, disabled = find_next(page)
        if nxt is None:
            job_log(job_dir, "No 'next page' button found: assuming this is the only/last page.", "WARNING")
            last = True
        else:
            _dispose(nxt)
            if disabled:
                job_log(job_dir, f"'Next' is disabled on page {n}: that was the last page.")
                last = True
    return last


def download_loop(page, job_dir: Path, dl_dir: Path, pages_dir: Path, rep: StepReporter, monitor, browser,
                  state: dict, start_d: date, end_d: date) -> str:
    """Run page cycles until the last page ('done') or until the browser should be restarted ('recycle')."""
    session_pages = 0
    baseline = None
    while True:
        n = state["resume_page"]
        last = False
        for attempt in range(1, CYCLE_RETRIES + 2):
            try:
                last = fetch_page(page, job_dir, dl_dir, pages_dir, rep, monitor, browser, state, n,
                                  start_d, end_d, first_cycle=(session_pages == 0 and attempt == 1))
                break
            except (PWError, RuntimeError) as exc:
                if monitor.tripped or not browser.is_connected() or page.is_closed() or attempt > CYCLE_RETRIES:
                    raise
                job_log(job_dir, f"P{n} cycle attempt {attempt}/{CYCLE_RETRIES + 1} failed: {str(exc)[:250]} "
                                 f"-> retrying the whole page cycle.", "WARNING")
                try:
                    page.keyboard.press("Escape")
                except PWError:
                    pass
                page.wait_for_timeout(2000)
        session_pages += 1

        snap = collect_snapshot(dl_dir)
        job_log(job_dir, f"P{n} DONE            | {format_snapshot(snap)}")
        if baseline is None:
            baseline = snap["chrom_rss"]
        if monitor.tripped:
            raise RuntimeError(monitor.tripped)
        if not browser.is_connected() or page.is_closed():
            raise RuntimeError("Chromium exited while downloading pages.")
        if last:
            return "done"
        if n >= MAX_PAGES:
            job_log(job_dir, f"Reached the safety cap of {MAX_PAGES} pages.", "WARNING")
            return "done"
        reason = recycle_reason(snap, baseline, session_pages)
        if reason:
            job_log(job_dir, f"RECYCLE requested after page {n}: {reason}. The browser will be restarted and "
                             f"resume at page {n + 1}.", "WARNING")
            return "recycle"


# =====================================================================
#  Compile all page files into one workbook
# =====================================================================
def _iter_xlsx(path: Path):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        for row in wb.worksheets[0].iter_rows(values_only=True):
            yield list(row)
    finally:
        wb.close()


def _iter_csv(path: Path):
    raw = path.read_bytes()
    text = None
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    yield from csv.reader(io.StringIO(text), dialect)


def iter_file_rows(path: Path):
    ext = path.suffix.lower()
    if ext == ".xls":
        raise RuntimeError("The site returned an old .xls file, which this app cannot read. "
                           "Install xlrd or ask me to add .xls support.")
    if ext in (".csv", ".txt"):
        return _iter_csv(path)
    return _iter_xlsx(path)


def split_header(rows_iter):
    """First row with at least 3 filled cells is the header. Returns (header_list, rest_iterator)."""
    for row in rows_iter:
        vals = ["" if c is None else str(c).strip() for c in row]
        if sum(1 for v in vals if v) >= 3:
            return vals, rows_iter
    return None, rows_iter


def unique_names(header):
    seen, out = {}, []
    for i, h in enumerate(header):
        h = h or f"col_{i + 1}"
        n = seen.get(h, 0)
        seen[h] = n + 1
        out.append(h if n == 0 else f"{h}_{n + 1}")
    return out


def compile_files(files, out_path: Path, job_dir: Path, rep: StepReporter) -> int:
    """Merge every page file into one XLSX (columns matched by header name). Streaming: no DataFrames. Returns rows."""
    master, seen, headers = [], set(), []
    for f in files:                                   # pass 1: collect the columns
        gen = iter_file_rows(f)
        header, _ = split_header(gen)
        if hasattr(gen, "close"):
            gen.close()
        if header is None:
            raise RuntimeError(f"Could not find a header row in {f.name}.")
        names = unique_names(header)
        headers.append(names)
        for n in names:
            if n not in seen:
                seen.add(n)
                master.append(n)

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Closed Alarm")
    ws.freeze_panes = "A2"
    for i, label in enumerate(master, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(len(label) + 4, 14), 40)
    ws.append(master)

    total = 0
    for idx, f in enumerate(files):                   # pass 2: write the rows
        names = headers[idx]
        pos = {n: i for i, n in enumerate(names)}
        order = [pos.get(m) for m in master]
        gen = iter_file_rows(f)
        _, rest = split_header(gen)
        count = 0
        for row in rest:
            if not any(c not in (None, "") for c in row):
                continue
            out = []
            for p in order:
                v = row[p] if p is not None and p < len(row) else None
                out.append(ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v)
            ws.append(out)
            count += 1
        if hasattr(gen, "close"):
            gen.close()
        total += count
        if idx < 3 or (idx + 1) % 20 == 0 or idx + 1 == len(files):
            job_log(job_dir, f"Compiled {f.name}: {count:,} rows (total {total:,}) | "
                             f"{format_snapshot(collect_snapshot(job_dir))}")
        rep.running(f"Compiling files into one Excel: {idx + 1}/{len(files)} done, {total:,} rows so far...",
                    0.90 + 0.08 * (idx + 1) / len(files), quiet=True)
    wb.save(out_path)
    return total


# =====================================================================
#  One browser SESSION: login -> CLOSED ALARM tab -> page cycles
#  (500 rows -> Start Date via calendar -> Apply -> verify 500 -> page N -> verify 500 -> download)
# =====================================================================
def run_session(job_dir: Path, dl_dir: Path, pages_dir: Path, rep: StepReporter, monitor: DiagnosticsMonitor,
                state: dict, username: str, password: str, start_d: date, end_d: date) -> str:
    first = state["session_no"] == 1
    resume_page = state["resume_page"]

    shutil.rmtree(dl_dir, ignore_errors=True)   # Playwright's temp download folder (page files live in pages_dir)
    dl_dir.mkdir(parents=True, exist_ok=True)
    monitor.reset()
    monitor.phase = "starting"

    job_log(job_dir, f"=== BROWSER SESSION {state['session_no']} starting (resume at page {resume_page}) | "
                     f"{format_snapshot(collect_snapshot(dl_dir))}")
    rep.running("Starting browser..." if first else
                f"Restarting the browser (session {state['session_no']}), resuming at page {resume_page}...",
                0.05 if first else 0.46)

    with sync_playwright() as pw:
        browser = launch_browser(pw, job_dir, dl_dir)
        page = None
        context = None
        try:
            context = browser.new_context(accept_downloads=True, viewport={"width": 1600, "height": 1000})
            if BLOCK_IMAGES:
                context.route(
                    "**/*",
                    lambda route: route.abort()
                    if route.request.resource_type in ("image", "media", "font")
                    else route.continue_(),
                )
            page = context.new_page()
            page.set_default_timeout(DEFAULT_TIMEOUT * 1000)
            attach_diagnostics(job_dir, browser, page)

            monitor.phase = "login"
            rep.running("Logging in...", 0.10, quiet=not first)
            pw_login(page, username, password)

            monitor.phase = "alarm-page"
            rep.running("Opening alarm page...", 0.15, quiet=not first)
            pw_open_alarm_page(page)
            pw_scroll_to_alarm_section(page)

            rep.running("Opening CLOSED ALARM tab...", 0.20, quiet=not first)
            pw_click_closed_alarm_tab(page, job_dir)
            if first:
                take_shot(page, job_dir, "shot_0_closed_alarm_tab")

            monitor.phase = "footer-wait"
            rep.running("Waiting for the table footer (pagination) to render...", 0.24, quiet=not first)
            wait_footer_ready(page, job_dir, timeout=60)

            monitor.phase = "download-pages"
            rep.running("Starting the page cycles (500 rows → start date → Apply → verify → page → download)...",
                        0.30 if first else 0.46)
            outcome = download_loop(page, job_dir, dl_dir, pages_dir, rep, monitor, browser, state, start_d, end_d)
            if outcome == "done" and first:
                take_shot(page, job_dir, "shot_4_last_page")
            return outcome
        except Exception:
            job_log(job_dir, "State at failure: " + format_snapshot(collect_snapshot(dl_dir)), "ERROR")
            if page is not None:
                try:
                    page.screenshot(path=str(job_dir / "error.png"), timeout=10000)
                except Exception:
                    pass
            raise
        finally:
            monitor.phase = "closing"
            SITE_INFO["quiet"] = False
            for closer in (context, browser):
                try:
                    if closer is not None:
                        closer.close()
                except Exception:
                    pass


# =====================================================================
#  STEP 1 worker: runs sessions (restarting the browser when needed), then compiles
# =====================================================================
def retrieve_worker(job_dir: Path, username: str, password: str, start_d: date, end_d: date):
    rep = StepReporter(job_dir, 1)
    dl_dir = job_dir / "dl"
    pages_dir = job_dir / "pages"
    shutil.rmtree(pages_dir, ignore_errors=True)
    pages_dir.mkdir(parents=True, exist_ok=True)
    dl_dir.mkdir(parents=True, exist_ok=True)
    for pattern in ("shot_*.png", "shot_*.jpg", "error.png", "dom_*.txt", "result_*.xlsx", "result.part"):
        for old in job_dir.glob(pattern):
            old.unlink(missing_ok=True)

    state = {"files": [], "hashes": {}, "resume_page": 1, "total_pages": 1, "session_no": 0,
             "page_size": TARGET_PAGE_SIZE, "size_discovered": False}
    monitor = None
    try:
        log_system_info(job_dir)
        job_log(job_dir, f"Requested Start Date: {start_d:%d-%m-%Y} (selected via the calendar). "
                         f"End Date is left to the site default (today); the form value {end_d:%d-%m-%Y} "
                         f"is only used for the result file name.")
        monitor = DiagnosticsMonitor(job_dir, dl_dir, rep)
        monitor.start()

        recycles = crash_restarts = 0
        while True:
            state["session_no"] += 1
            resume_before = state["resume_page"]
            try:
                outcome = run_session(job_dir, dl_dir, pages_dir, rep, monitor, state, username, password,
                                      start_d, end_d)
            except Exception as exc:
                reason = monitor.tripped or (str(exc).strip().splitlines()[0] if str(exc).strip() else type(exc).__name__)
                progressed = state["resume_page"] > resume_before
                if progressed and crash_restarts < MAX_CRASH_RESTARTS:
                    crash_restarts += 1
                    job_log(job_dir, f"Session {state['session_no']} failed after progress "
                                     f"(pages saved up to {state['resume_page'] - 1}): {reason[:250]} -> "
                                     f"restart {crash_restarts}/{MAX_CRASH_RESTARTS}, resuming at page "
                                     f"{state['resume_page']}.", "ERROR")
                    kill_browser_processes(job_dir)
                    time.sleep(2)
                    continue
                if monitor.tripped:
                    raise RuntimeError(monitor.tripped)
                raise
            if outcome == "recycle":
                recycles += 1
                if recycles > MAX_RECYCLES:
                    raise RuntimeError(f"Browser was recycled {MAX_RECYCLES} times; stopping at page "
                                       f"{state['resume_page']}.")
                monitor.phase = "recycling"
                job_log(job_dir, f"Browser recycle {recycles}/{MAX_RECYCLES} done; next page "
                                 f"{state['resume_page']}. Memory after close: "
                                 f"{format_snapshot(collect_snapshot(dl_dir))}")
                time.sleep(2)
                continue
            break

        files = sorted(state["files"])
        if not files:
            raise RuntimeError("No page files were downloaded.")
        monitor.phase = "compile"
        rep.running(f"Compiling {len(files)} downloaded files into one Excel...", 0.90)
        result_name = f"closed_alarm_{start_d:%Y%m%d}_{end_d:%Y%m%d}.xlsx"
        part = job_dir / "result.part"
        total_rows = compile_files(files, part, job_dir, rep)
        os.replace(part, job_dir / result_name)
        size_mb = (job_dir / result_name).stat().st_size / 1024 / 1024
        rep.done(f"Compiled {total_rows:,} rows from {len(files)} page files "
                 f"({start_d:%d-%m-%Y} to {end_d:%d-%m-%Y}, {size_mb:.1f} MB).",
                 result_file=result_name, rows=total_rows, pages=len(files),
                 range_start=start_d.isoformat(), range_end=end_d.isoformat())
    except Exception as exc:
        job_log(job_dir, traceback.format_exc(), "ERROR")
        if monitor is not None and monitor.tripped:
            message = monitor.tripped
        else:
            message = str(exc).strip().splitlines()[0] if str(exc).strip() else type(exc).__name__
        rep.fail(message)
    finally:
        if monitor is not None:
            monitor.phase = "finished"
            monitor.stop()
            job_log(job_dir, f"SUMMARY: peak working set={monitor.peak_used:.0f}MB, "
                             f"peak Chromium RSS={monitor.peak_chromium:.0f}MB, "
                             f"browser sessions={state['session_no']}, "
                             f"memory guard tripped={'YES' if monitor.tripped else 'no'}")
        shutil.rmtree(dl_dir, ignore_errors=True)
        shutil.rmtree(pages_dir, ignore_errors=True)
        (job_dir / "result.part").unlink(missing_ok=True)


# =====================================================================
#  UI
# =====================================================================
def start_thread(target, *args):
    threading.Thread(target=target, args=args, daemon=True).start()


def get_job_id() -> str:
    jid = st.query_params.get("job", "")
    if not re.fullmatch(r"[a-f0-9]{12}", jid or ""):
        jid = uuid.uuid4().hex[:12]
        st.query_params["job"] = jid
    return jid


def fmt_elapsed(started):
    secs = int(time.time() - started)
    return f"{secs // 60}m {secs % 60:02d}s"


def show_status(status: dict):
    state = status.get("state", "idle")
    msg = status.get("message", "")
    if state == "running":
        if status.get("progress") is not None:
            st.progress(min(max(status["progress"], 0.0), 1.0))
        st.info(f"⏳ {msg}  \n_Total time: {fmt_elapsed(status.get('started', time.time()))}_")
    elif state == "done":
        st.success(f"✅ {msg}")
    elif state == "error":
        st.error(f"❌ {msg}")


def default_range():
    """Last two months: the 1st of the previous month up to today (site time)."""
    today = datetime.now(WIB).date()
    prev_month_last_day = today.replace(day=1) - timedelta(days=1)
    return prev_month_last_day.replace(day=1), today


def read_log_tail(log_file: Path, n: int):
    """Last n lines of the log without loading a huge file into memory."""
    try:
        size = log_file.stat().st_size
        with open(log_file, "rb") as fh:
            fh.seek(max(0, size - 64 * 1024))
            data = fh.read().decode("utf-8", errors="replace")
        return data.splitlines()[-n:]
    except OSError:
        return []


def render_workflow(job_dir: Path, job_id: str):
    s1 = get_status(job_dir, 1)
    busy = s1["state"] == "running"

    was_busy = st.session_state.get("_was_busy", False)
    st.session_state["_was_busy"] = busy
    if was_busy and not busy:
        st.rerun()

    result_path = job_dir / s1["result_file"] if s1.get("state") == "done" and s1.get("result_file") else None
    result_ok = result_path is not None and result_path.exists()

    # ------------------------------------------------------------ STEP 1
    st.subheader("Step 1 · Retrieve closed alarms")
    st.caption("Make sure to select the correct Start Date and ensure the date range does not exceed 30 days. "
               "considering that the End Date is automatically set to the CURRENT DAY. "
               )
    show_status(s1)
    if s1["state"] == "error" and (job_dir / "error.png").exists():
        st.image(str(job_dir / "error.png"), caption="Browser screenshot when the error happened")

    d0, d1 = default_range()
    with st.form("step1_form"):
        username = st.text_input("Username", disabled=busy)
        password = st.text_input("Password", type="password", disabled=busy)
        cs, ce = st.columns(2)
        start_d = cs.date_input("Start date (selected in the site calendar)", value=d0, format="DD/MM/YYYY",
                                disabled=busy, key="rng_start")
        end_d = ce.date_input("End date (site default = today, not changed)", value=d1, format="DD/MM/YYYY",
                              disabled=True, key="rng_end")
        force = st.checkbox("Replace the existing result (retrieve again)", disabled=busy or not result_ok)
        go1 = st.form_submit_button("Retrieve closed alarms", type="primary", disabled=busy, use_container_width=True)

    if go1:
        if result_ok and not force:
            st.info("A compiled result already exists, so retrieving was skipped. "
                    "Tick “Replace the existing result” to retrieve a new copy.")
        elif not username or not password:
            st.error("Please enter both username and password.")
        elif start_d > end_d:
            st.error("Start date must not be after today's date (the site's End Date).")
        else:
            (job_dir / "error.png").unlink(missing_ok=True)
            for old in job_dir.glob("closed_alarm_*.xlsx"):
                old.unlink(missing_ok=True)
            StepReporter(job_dir, 1).running("Queued...", 0.0)
            start_thread(retrieve_worker, job_dir, username, password, start_d, end_d)
            st.session_state["_was_busy"] = True
            st.rerun()

    st.divider()

    # ------------------------------------------------------------ STEP 2
    st.subheader("Step 2 · Download Result")
    st.caption("Just hands you the compiled file from Step 1. No processing happens here.")
    if result_ok:
        rs, re_ = s1.get("range_start"), s1.get("range_end")
        size_txt = f"{result_path.stat().st_size / 1024 / 1024:.2f} MB"
        st.download_button(
            "Download Result",
            data=result_path.read_bytes(),
            file_name=result_path.name,
            mime=XLSX_MIME,
            type="primary",
            use_container_width=True,
            on_click="ignore",
        )
        if isinstance(s1.get("rows"), int):
            txt = f"{s1['rows']:,} rows · {s1.get('pages', '?')} pages · {size_txt}"
            if rs and re_:
                txt += f" · {date.fromisoformat(rs):%d-%m-%Y} to {date.fromisoformat(re_):%d-%m-%Y}"
        else:
            txt = size_txt
        st.caption(txt)
    else:
        st.button("Download Result", disabled=True, use_container_width=True, key="dl_disabled")
        st.caption("🔒 Complete Step 1 first.")

    # ------------------------------------------------------------ extras
    st.divider()
    with st.expander("Diagnostics (log, memory, screenshots)"):
        log_file = job_dir / "job.log"
        if busy:
            # keep the UI light while the worker runs: tail only, no screenshots / download buttons
            st.code("\n".join(read_log_tail(log_file, 40)) or "(empty)", language=None)
            st.caption("Full log, DOM dumps and screenshots appear here when the step finishes.")
        else:
            all_lines = log_file.read_text(encoding="utf-8").splitlines() if log_file.exists() else []
            st.code("\n".join(all_lines[-150:]) or "(empty)", language=None)
            if all_lines:
                st.download_button("Download full log (job.log)", data="\n".join(all_lines),
                                   file_name=f"job_{job_id}.log", mime="text/plain", on_click="ignore", key="dl_log")
            for dom in sorted(job_dir.glob("dom_*.txt")):
                st.download_button(f"Download {dom.name} (DOM dump)", data=dom.read_bytes(),
                                   file_name=dom.name, mime="text/plain", on_click="ignore", key=f"dl_{dom.name}")
            for shot in sorted(job_dir.glob("shot_*.jpg")):
                st.image(str(shot), caption=shot.stem)
    if st.button("🗑️ Delete my files & start over", disabled=busy):
        shutil.rmtree(job_dir, ignore_errors=True)
        st.query_params.clear()
        st.session_state.clear()
        st.rerun()


def main():
    st.set_page_config(page_title="PMT Alarm Downloader", page_icon="🚨", layout="centered")
    st.title("🚨 PMT Alarm – Automate Extraction")
    cleanup_old_jobs()

    job_id = get_job_id()
    job_dir = JOBS_ROOT / job_id
    job_dir.mkdir(parents=True, exist_ok=True)
    st.caption(
        f"Session `{job_id}` – the page URL contains this id. If the page reloads or disconnects "
        "during a long step, reopen the same URL to continue."
    )

    busy = get_status(job_dir, 1)["state"] == "running"
    st.fragment(run_every=POLL_SECONDS if busy else None)(render_workflow)(job_dir, job_id)

    st.caption("Your credentials are used only to start Step 1 and are never written to disk or logs.")


if __name__ == "__main__":
    main()
